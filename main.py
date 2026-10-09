#!/usr/bin/env python3
"""
Automation Hub launcher (pywebview).

One window, one shared base folder ("ROOT"). The hub authenticates the user
against ROOT/master/user.csv, lists the HTML tools in ROOT/apps/ that the user
may open, and exposes a filesystem bridge (LauncherAPI) to those tools as
``window.pywebview.api``. Every relative path handed to the bridge is resolved
inside ROOT and cannot escape it.

The base folder is chosen once and stored in launcher_config.json next to this
file. Changing it later is admin-password gated (ROOT/pwd.txt, default
"admin123") and the hub can either *switch* to another existing folder or
*migrate* every file of the current folder to a new location (copy, verify,
switch, optionally remove the originals). See folder_manager.py.

Run:  python main.py [--debug]
"""

from __future__ import annotations

import base64
import contextlib
import csv
import datetime as _dt
import getpass
import hashlib
import hmac
import io
import json
import logging
import os
import re
import subprocess
import sys
import threading
import time
import uuid
from collections import defaultdict
from pathlib import Path

import folder_manager as fm

__version__ = "2.0.0"

APP_DIR = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent
CONFIG_PATH = APP_DIR / "launcher_config.json"
STORAGE_DIR = APP_DIR / "webview_storage"
DEFAULT_ADMIN_PASSWORD = "admin123"
LEGACY_FOLDER_KEYS = ("folder_path", "folder", "base_folder", "root", "path")
ACCESS_TRUE = {"1", "y", "yes", "true", "x"}

log = logging.getLogger("hub")


# ──────────────────────────────────────────────────────────────────────────────
# Small helpers
# ──────────────────────────────────────────────────────────────────────────────

def _os_user() -> str:
    try:
        name = getpass.getuser()
    except Exception:
        name = os.environ.get("USERNAME") or os.environ.get("USER") or ""
    return name.split("\\")[-1].strip()


def _safe_segment(value: str, max_len: int = 80) -> str:
    """Make a single path segment safe on Windows and POSIX."""
    s = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(value or "")).strip(" .")
    return (s or "_")[:max_len]


def _decode(raw: bytes) -> str:
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return str(int(v))
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    return str(v)


def _dialog_type(kind: str):
    """webview.FileDialog.<KIND> on pywebview >= 5, FOLDER_DIALOG/SAVE_DIALOG before."""
    import webview
    fd = getattr(webview, "FileDialog", None)
    if fd is not None and hasattr(fd, kind):
        return getattr(fd, kind)
    return getattr(webview, f"{kind}_DIALOG")


def _hash_pwd(pwd: str) -> str:
    return hashlib.sha256(("hub:" + (pwd or "")).encode("utf-8")).hexdigest()


def _atomic_write_bytes(path: Path, data: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.name}.{uuid.uuid4().hex[:8]}.tmp")
    try:
        with open(tmp, "wb") as fh:
            fh.write(data)
            fh.flush()
            with contextlib.suppress(OSError):
                os.fsync(fh.fileno())
        for attempt in range(6):
            try:
                os.replace(tmp, path)
                return
            except PermissionError:
                # Windows: another process (AV scanner, Excel, a colleague) has
                # the target open for a moment. Retry briefly.
                if attempt == 5:
                    raise
                time.sleep(0.15 * (attempt + 1))
    finally:
        if tmp.exists():
            with contextlib.suppress(OSError):
                tmp.unlink()


def _atomic_write_text(path: Path, text: str) -> None:
    _atomic_write_bytes(path, (text or "").encode("utf-8"))


def load_config() -> dict:
    try:
        with open(CONFIG_PATH, "r", encoding="utf-8-sig") as fh:
            cfg = json.load(fh)
        return cfg if isinstance(cfg, dict) else {}
    except (OSError, ValueError):
        return {}


def save_config(cfg: dict) -> None:
    _atomic_write_text(CONFIG_PATH, json.dumps(cfg, indent=2, ensure_ascii=False))


def config_folder(cfg: dict) -> str:
    for k in LEGACY_FOLDER_KEYS:
        v = cfg.get(k)
        if isinstance(v, str) and v.strip():
            return v.strip()
    return ""


# ──────────────────────────────────────────────────────────────────────────────
# Bridge
# ──────────────────────────────────────────────────────────────────────────────

class LauncherAPI:
    """Exposed to every page as window.pywebview.api.

    Methods starting with an underscore are not exposed by pywebview.
    """

    def __init__(self, config_path: Path | None = None):
        self._window = None
        self._config_path = Path(config_path) if config_path else None
        self.current_user: dict | None = None
        self.folder_path: str | None = None
        self.current_app: str | None = None
        self._path_locks: dict[str, threading.Lock] = defaultdict(threading.Lock)
        self._locks_guard = threading.Lock()
        self._inflight = 0
        self._inflight_lock = threading.Lock()
        self._migration: fm.MigrationJob | None = None
        self._auto_login = True
        self._xl = None  # excel_link.ExcelLinkEngine, created on first use
        cfg = self._load_cfg()
        folder = config_folder(cfg)
        self.folder_path = fm.norm(folder) if folder else None

    # Some older call sites referenced self.data_root (it was never set, which
    # broke attachments). Keep it as an alias of folder_path.
    @property
    def data_root(self) -> str | None:
        return self.folder_path

    # ---- config --------------------------------------------------------------
    def _load_cfg(self) -> dict:
        if self._config_path:
            try:
                return json.loads(self._config_path.read_text(encoding="utf-8-sig"))
            except (OSError, ValueError):
                return {}
        return load_config()

    def _save_cfg(self, cfg: dict) -> None:
        if self._config_path:
            _atomic_write_text(self._config_path, json.dumps(cfg, indent=2, ensure_ascii=False))
        else:
            save_config(cfg)

    def _remember_folder(self, folder: str) -> None:
        cfg = self._load_cfg()
        old = config_folder(cfg)
        for k in LEGACY_FOLDER_KEYS:
            cfg.pop(k, None)
        cfg["folder_path"] = folder
        recent = [f for f in cfg.get("recent_folders", []) if not fm.same_path(f, folder)]
        if old and not fm.same_path(old, folder):
            recent = [f for f in recent if not fm.same_path(f, old)]
            recent.insert(0, old)
        cfg["recent_folders"] = recent[:8]
        cfg["updated_at"] = fm._now_iso()
        self._save_cfg(cfg)
        self.folder_path = folder

    # ---- window --------------------------------------------------------------
    def _set_window(self, window) -> None:
        self._window = window

    def _on_closing(self):
        # Give in-flight bridge writes a moment to land before the process exits.
        deadline = time.time() + 2.0
        while time.time() < deadline:
            with self._inflight_lock:
                if self._inflight <= 0:
                    break
            time.sleep(0.05)
        return True

    @contextlib.contextmanager
    def _writing(self, path: Path):
        key = os.path.normcase(str(path))
        with self._locks_guard:
            lock = self._path_locks[key]
        with self._inflight_lock:
            self._inflight += 1
        try:
            with lock:
                yield
        finally:
            with self._inflight_lock:
                self._inflight -= 1

    # ---- paths ---------------------------------------------------------------
    def _root(self) -> Path:
        if not self.folder_path:
            raise RuntimeError("The base folder is not set.")
        return Path(self.folder_path)

    def _abs(self, rel) -> Path:
        """Resolve a relative path inside ROOT; refuse anything that escapes it."""
        root = self._root()
        r = str(rel or "").replace("\\", "/").strip()
        if re.match(r"^[A-Za-z]:", r) or r.startswith("//") or r.startswith("/"):
            # Absolute paths are accepted only when they already point inside
            # ROOT (legacy attachment records stored absolute paths).
            cand = fm.norm(r)
            if fm.same_path(cand, str(root)) or fm.is_inside(cand, str(root)):
                return Path(cand)
            if re.match(r"^[A-Za-z]:", r) or r.startswith("//"):
                raise ValueError(f"Path is outside the base folder: {rel}")
            r = r.lstrip("/")  # "/reco/x.json" means ROOT-relative
        cand = os.path.normpath(os.path.join(str(root), *[p for p in r.split("/") if p]))
        if not (fm.same_path(cand, str(root)) or fm.is_inside(cand, str(root))):
            raise ValueError(f"Path escapes the base folder: {rel}")
        return Path(cand)

    def _rel(self, path: Path) -> str:
        return os.path.relpath(str(path), str(self._root())).replace("\\", "/")

    def _read_text(self, path: Path) -> str:
        with open(path, "rb") as fh:
            return _decode(fh.read())

    # ──────────────────────────────────────────────────────────────────────
    # Hub: state, users, apps
    # ──────────────────────────────────────────────────────────────────────
    def init_app(self) -> dict:
        """Everything the hub page needs to render. Auto-signs-in the OS user."""
        cfg = self._load_cfg()
        state = {
            "version": __version__,
            "os_user": _os_user(),
            "configured": bool(self.folder_path),
            "folder": self.folder_path or "",
            "folder_ok": False,
            "folder_info": None,
            "pointer": None,
            "users_file": False,
            "user": None,
            "login_error": "",
            "apps": [],
            "hidden_apps": 0,
            "recent_folders": cfg.get("recent_folders", []),
            "migration": self._migration.to_dict() if self._migration else None,
        }
        if not self.folder_path:
            return state
        info = fm.inspect_folder(self.folder_path)
        state["folder_info"] = info
        state["folder_ok"] = info["exists"] and info["is_dir"]
        if not state["folder_ok"]:
            return state
        state["pointer"] = info.get("pointer")
        users = self._read_users()
        state["users_file"] = users is not None
        if self.current_user is None and users and self._auto_login:
            match = self._match_user(users, _os_user())
            if match:
                self.current_user = match
        if self.current_user:
            state["user"] = self._public_user(self.current_user)
            state["apps"], state["hidden_apps"] = self._apps_for(self.current_user)
        return state

    def login(self, user_id: str) -> dict:
        users = self._read_users()
        if users is None:
            return {"ok": False, "error": "master/user.csv was not found in the base folder."}
        match = self._match_user(users, user_id)
        if not match:
            return {"ok": False, "error": f"'{user_id}' is not listed in master/user.csv."}
        self.current_user = match
        return {"ok": True, "state": self.init_app()}

    def logout(self) -> dict:
        # After an explicit "Switch user" do not silently sign the OS user back in.
        self.current_user = None
        self._auto_login = False
        return self.init_app()

    def current_username(self) -> str:
        # Apps can ask before the hub page has run init_app (e.g. after a reload).
        if not self.current_user and self._auto_login and self.folder_path:
            users = self._read_users()
            if users:
                self.current_user = self._match_user(users, _os_user())
        if self.current_user:
            return self.current_user.get("name") or self.current_user.get("id") or _os_user()
        return _os_user()

    def _read_users(self) -> list[dict] | None:
        try:
            p = self._abs("master/user.csv")
        except (RuntimeError, ValueError):
            return None
        if not p.is_file():
            return None
        try:
            rows = list(csv.reader(io.StringIO(self._read_text(p))))
        except OSError:
            return None
        if not rows:
            return []
        keys = [h.strip().lower() for h in rows[0]]
        users = []
        for r in rows[1:]:
            if not r or not r[0].strip():
                continue
            access = {}
            for i in range(3, len(keys)):
                if keys[i]:
                    access[keys[i]] = (r[i].strip().lower() if i < len(r) else "") in ACCESS_TRUE
            users.append({
                "id": r[0].strip(),
                "name": r[1].strip() if len(r) > 1 else "",
                "email": r[2].strip() if len(r) > 2 else "",
                "access": access,
            })
        return users

    @staticmethod
    def _match_user(users: list[dict], who: str) -> dict | None:
        w = str(who or "").strip().split("\\")[-1].lower()
        if not w:
            return None
        for u in users:
            uid = u["id"].split("\\")[-1].lower()
            if uid == w or u["email"].lower() == w or u["email"].lower().split("@")[0] == w:
                return u
        return None

    @staticmethod
    def _public_user(u: dict) -> dict:
        return {"id": u["id"], "name": u["name"] or u["id"], "email": u["email"]}

    def _app_files(self) -> list[tuple[str, Path]]:
        try:
            apps_dir = self._abs("apps")
        except (RuntimeError, ValueError):
            return []
        if not apps_dir.is_dir():
            return []
        out = []
        for p in sorted(apps_dir.iterdir(), key=lambda x: x.name.lower()):
            if p.is_file() and p.suffix.lower() in (".html", ".htm"):
                out.append((p.stem.lower(), p))
        return out

    def _icon_for(self, key: str) -> str:
        for rel in (f"apps/icon/{key}", f"icon/{key}"):
            for ext, mime in ((".png", "image/png"), (".svg", "image/svg+xml"),
                              (".jpg", "image/jpeg"), (".ico", "image/x-icon")):
                try:
                    p = self._abs(rel + ext)
                    if p.is_file() and p.stat().st_size < 2_000_000:
                        return f"data:{mime};base64," + base64.b64encode(p.read_bytes()).decode("ascii")
                except (OSError, ValueError, RuntimeError):
                    continue
        return ""

    def _apps_for(self, user: dict) -> tuple[list[dict], int]:
        apps, hidden = [], 0
        for key, path in self._app_files():
            if user["access"].get(key):
                apps.append({"key": key, "name": path.stem, "file": self._rel(path), "icon": self._icon_for(key)})
            else:
                hidden += 1
        return apps, hidden

    def list_apps(self) -> list[dict]:
        if not self.current_user:
            return []
        return self._apps_for(self.current_user)[0]

    def open_app(self, key: str) -> dict:
        if not self.current_user:
            return {"ok": False, "error": "Sign in first."}
        for k, path in self._app_files():
            if k == str(key).lower():
                if not self.current_user["access"].get(k):
                    return {"ok": False, "error": "You do not have access to this app."}
                self.current_app = k
                if self._window:
                    with contextlib.suppress(Exception):
                        self._window.set_title(f"Automation Hub — {path.stem}")
                    self._window.load_url(path.as_uri())
                return {"ok": True, "url": path.as_uri()}
        return {"ok": False, "error": f"App '{key}' not found in apps/."}

    def return_to_hub(self) -> bool:
        self.current_app = None
        if self._window:
            with contextlib.suppress(Exception):
                self._window.set_title("Automation Hub")
            self._window.load_html(HUB_HTML)
        return True

    def close_window(self) -> bool:
        if self._window:
            self._window.destroy()
        return True

    # ──────────────────────────────────────────────────────────────────────
    # Base folder: admin, switch, migrate
    # ──────────────────────────────────────────────────────────────────────
    def _expected_admin_password(self) -> str | None:
        """pwd.txt of the current base folder; None when it is unreachable."""
        if not self.folder_path or not os.path.isdir(self.folder_path):
            return None
        p = Path(self.folder_path) / "pwd.txt"
        try:
            if p.is_file():
                return self._read_text(p).strip()
        except OSError:
            return None
        return DEFAULT_ADMIN_PASSWORD

    def _check_admin(self, pwd: str) -> bool:
        pwd = str(pwd or "")
        if not self.folder_path:
            return True  # first-time setup: nothing to protect yet
        expected = self._expected_admin_password()
        cfg = self._load_cfg()
        if expected is not None:
            ok = hmac.compare_digest(pwd, expected)
            if ok and cfg.get("admin_hash") != _hash_pwd(pwd):
                cfg["admin_hash"] = _hash_pwd(pwd)
                with contextlib.suppress(OSError):
                    self._save_cfg(cfg)
            return ok
        # Base folder offline: fall back to the last verified password.
        cached = cfg.get("admin_hash")
        if cached:
            return hmac.compare_digest(_hash_pwd(pwd), cached)
        return hmac.compare_digest(pwd, DEFAULT_ADMIN_PASSWORD)

    def verify_admin_password(self, pwd: str) -> bool:
        return self._check_admin(pwd)

    def browse_folder(self, start: str = "") -> str:
        if not self._window:
            return ""
        directory = start if start and os.path.isdir(start) else (self.folder_path or "")
        try:
            res = self._window.create_file_dialog(_dialog_type("FOLDER"), directory=directory or "")
        except Exception as exc:
            log.warning("folder dialog failed: %s", exc)
            return ""
        if not res:
            return ""
        return fm.norm(res[0] if isinstance(res, (list, tuple)) else res)

    def inspect_path(self, path: str) -> dict:
        info = fm.inspect_folder(path)
        cur = self.folder_path or ""
        info["is_current"] = bool(cur and path and fm.same_path(path, cur))
        info["inside_current"] = bool(cur and path and fm.is_inside(path, cur))
        info["contains_current"] = bool(cur and path and fm.is_inside(cur, path))
        return info

    def get_folder_info(self) -> dict:
        info = fm.inspect_folder(self.folder_path) if self.folder_path else None
        return {"folder": self.folder_path or "", "info": info,
                "recent_folders": self._load_cfg().get("recent_folders", [])}

    def set_base_folder(self, path: str, pwd: str = "", initialize: bool = False) -> dict:
        """Point the hub at another *existing* folder (no files are copied)."""
        if self._migration and not self._migration.finished:
            return {"ok": False, "error": "A migration is running. Wait for it to finish or cancel it."}
        if not self._check_admin(pwd):
            return {"ok": False, "error": "Incorrect admin password."}
        info = fm.inspect_folder(path)
        if not (info["exists"] and info["is_dir"]):
            return {"ok": False, "error": info["error"] or "Folder not found."}
        if not info["writable"]:
            return {"ok": False, "error": "This folder is not writable for your account."}
        if not info["looks_like_root"] and not initialize:
            return {"ok": False, "needs_init": True, "info": info,
                    "error": "This folder does not contain apps/, master/ or reco/."}
        created = []
        if initialize:
            created = fm.init_skeleton(info["path"], _os_user(), _os_user(),
                                       app_keys=self._app_keys_in(info["path"]) or ["reco"])
        self._remember_folder(info["path"])
        self.current_user = None
        self._auto_login = True
        return {"ok": True, "created": created, "state": self.init_app()}

    def initialize_folder(self, pwd: str = "") -> dict:
        """Create the standard layout (and master/user.csv) in the current folder."""
        if not self._check_admin(pwd):
            return {"ok": False, "error": "Incorrect admin password."}
        root = str(self._root())
        created = fm.init_skeleton(root, _os_user(), _os_user(),
                                   app_keys=self._app_keys_in(root) or ["reco"])
        self.current_user = None
        return {"ok": True, "created": created, "state": self.init_app()}

    @staticmethod
    def _app_keys_in(root: str) -> list[str]:
        apps = Path(root) / "apps"
        if not apps.is_dir():
            return []
        return sorted({p.stem.lower() for p in apps.iterdir()
                       if p.is_file() and p.suffix.lower() in (".html", ".htm")})

    def follow_pointer(self) -> dict:
        """Switch to the folder named in MIGRATED_TO.txt (written by a migration)."""
        if not self.folder_path:
            return {"ok": False, "error": "No base folder set."}
        ptr = fm.read_pointer(self.folder_path)
        if not ptr:
            return {"ok": False, "error": "No migration note found in this folder."}
        info = fm.inspect_folder(ptr["new_folder"])
        if not (info["exists"] and info["is_dir"]):
            return {"ok": False, "error": f"The new folder is not reachable: {ptr['new_folder']}"}
        self._remember_folder(info["path"])
        self.current_user = None
        return {"ok": True, "state": self.init_app()}

    def preflight_migration(self, dest: str, pwd: str = "", conflict: str = "abort") -> dict:
        if not self._check_admin(pwd):
            return {"ok": False, "errors": ["Incorrect admin password."], "warnings": []}
        if not self.folder_path:
            return {"ok": False, "errors": ["No current base folder to migrate."], "warnings": []}
        return fm.preflight(self.folder_path, dest, conflict)

    def start_migration(self, dest: str, pwd: str = "", options: dict | None = None) -> dict:
        if not self._check_admin(pwd):
            return {"ok": False, "error": "Incorrect admin password."}
        if not self.folder_path:
            return {"ok": False, "error": "No current base folder to migrate."}
        if self._migration and not self._migration.finished:
            return {"ok": False, "error": "A migration is already running."}
        options = options or {}
        conflict = options.get("conflict", "abort")
        pre = fm.preflight(self.folder_path, dest, conflict)
        if not pre["ok"]:
            return {"ok": False, "error": "; ".join(pre["errors"]), "preflight": pre}
        switch = options.get("switch", True)

        def on_success(job: fm.MigrationJob) -> None:
            if switch:
                self._remember_folder(job.dst)
                self.current_user = None

        self._migration = fm.MigrationJob(
            self.folder_path, dest, conflict=conflict,
            remove_source=bool(options.get("remove_source")),
            verify_hash=bool(options.get("verify_hash")),
            on_success=on_success, user=self.current_username(),
        ).start()
        return {"ok": True, "job": self._migration.to_dict()}

    def migration_status(self) -> dict | None:
        return self._migration.to_dict() if self._migration else None

    def cancel_migration(self) -> bool:
        if self._migration and not self._migration.finished:
            self._migration.cancel()
            return True
        return False

    # ──────────────────────────────────────────────────────────────────────
    # Filesystem bridge used by the apps (paths relative to ROOT)
    # ──────────────────────────────────────────────────────────────────────
    def read_file(self, rel: str) -> str:
        """File text, or "" when the file does not exist.

        Raises when the file exists but cannot be read, or when the base folder
        itself is unreachable, so the app can tell "missing" (start empty) from
        "unreadable" (do not overwrite it with an empty copy later).
        """
        p = self._abs(rel)
        if p.is_file():
            return self._read_text(p)
        if not os.path.isdir(self._root()):
            raise OSError(f"The base folder is not reachable: {self.folder_path}")
        return ""

    def write_file(self, rel: str, data) -> bool:
        p = self._abs(rel)
        text = data if isinstance(data, str) else json.dumps(data, indent=2, ensure_ascii=False)
        with self._writing(p):
            _atomic_write_text(p, text)
        return True

    def append_csv_rows(self, rel: str, fields: list | None, rows: list | None) -> int:
        """Append rows (dicts) to a CSV without rewriting it.

        Creates the file with a header when missing. When the rows bring
        columns the header lacks, the file is rewritten once with the union
        header so nothing is dropped.
        """
        p = self._abs(rel)
        fields = [str(f) for f in (fields or [])]
        rows = [r for r in (rows or []) if isinstance(r, dict)]
        wanted = list(dict.fromkeys(fields + [k for r in rows for k in r.keys()]))
        with self._writing(p):
            text = self._read_text(p) if p.is_file() and p.stat().st_size > 0 else ""
            header = next(csv.reader(io.StringIO(text)), None) if text else None
            if not header:
                buf = io.StringIO()
                w = csv.writer(buf, lineterminator="\r\n")
                w.writerow(wanted)
                for r in rows:
                    w.writerow([_cell(r.get(h, "")) for h in wanted])
                _atomic_write_text(p, buf.getvalue())
                return len(rows)
            missing = [f for f in wanted if f not in header]
            if missing:
                new_header = header + missing
                existing = list(csv.DictReader(io.StringIO(text)))
                buf = io.StringIO()
                w = csv.writer(buf, lineterminator="\r\n")
                w.writerow(new_header)
                for r in existing:
                    w.writerow([_cell(r.get(h, "")) for h in new_header])
                for r in rows:
                    w.writerow([_cell(r.get(h, "")) for h in new_header])
                _atomic_write_text(p, buf.getvalue())
                return len(rows)
            buf = io.StringIO()
            w = csv.writer(buf, lineterminator="\r\n")
            for r in rows:
                w.writerow([_cell(r.get(h, "")) for h in header])
            prefix = "" if text.endswith("\n") else "\r\n"
            with open(p, "a", encoding="utf-8", newline="") as fh:
                fh.write(prefix + buf.getvalue())
            return len(rows)

    def path_exists(self, rel: str) -> bool:
        try:
            return self._abs(rel).exists()
        except (ValueError, RuntimeError):
            return False

    def path_exists_json(self, rel: str) -> str:
        try:
            p = self._abs(rel)
            return json.dumps({"exists": p.exists(), "is_dir": p.is_dir(), "is_file": p.is_file()})
        except (ValueError, RuntimeError) as exc:
            return json.dumps({"exists": False, "is_dir": False, "is_file": False, "error": str(exc)})

    def resolve_url(self, rel: str) -> str:
        """file:// URL for a ROOT-relative file (also checks ROOT/apps/<rel>)."""
        for cand in (rel, "apps/" + str(rel or "").lstrip("/")):
            try:
                p = self._abs(cand)
            except (ValueError, RuntimeError):
                continue
            if p.is_file():
                return p.as_uri()
        return ""

    def list_directories(self, rel: str = "") -> list[str]:
        try:
            p = self._abs(rel)
            if not p.is_dir():
                return []
            return sorted((c.name for c in p.iterdir() if c.is_dir()), key=str.lower)
        except (OSError, ValueError, RuntimeError):
            return []

    def list_files(self, rel: str = "") -> list[dict]:
        """Files directly inside a folder: [{name, ext, size_kb, modified}]."""
        try:
            p = self._abs(rel)
            if not p.is_dir():
                return []
            out = []
            for c in sorted(p.iterdir(), key=lambda x: x.name.lower()):
                if c.is_file() and not c.name.startswith(".") and not c.name.endswith(".tmp"):
                    st = c.stat()
                    out.append({"name": c.name, "ext": c.suffix.lower().lstrip("."),
                                "size_kb": round(st.st_size / 1024, 1),
                                "modified": _dt.datetime.fromtimestamp(st.st_mtime).isoformat(timespec="seconds")})
            return out
        except (OSError, ValueError, RuntimeError):
            return []

    def list_folder_files(self, rel: str = "") -> str:
        """Legacy form of list_files(): the same list as a JSON string."""
        return json.dumps(self.list_files(rel))

    # ---- accounts master + AEG templates --------------------------------------
    def load_master(self) -> str:
        return self.read_file("master/accounts_master.csv")

    def save_master(self, csv_text: str) -> bool:
        return self.write_file("master/accounts_master.csv", csv_text or fm.ACCOUNTS_MASTER_HEADER)

    def load_templates(self) -> list[dict]:
        out = []
        try:
            d = self._abs("aegtemplates")
        except (ValueError, RuntimeError):
            return out
        if not d.is_dir():
            return out
        for p in sorted(d.glob("*.json"), key=lambda x: x.name.lower()):
            try:
                obj = json.loads(self._read_text(p))
                if isinstance(obj, dict):
                    out.append(obj)
            except (OSError, ValueError) as exc:
                log.warning("skipping bad AEG template %s: %s", p.name, exc)
        return out

    def save_template(self, template_id: str, json_text) -> bool:
        return self.write_file(f"aegtemplates/{_safe_segment(template_id)}.json", json_text)

    def delete_template(self, template_id: str) -> bool:
        p = self._abs(f"aegtemplates/{_safe_segment(template_id)}.json")
        with self._writing(p):
            if p.exists():
                p.unlink()
        return True

    # ---- attachments -----------------------------------------------------------
    def save_attachment(self, template_id: str, filename: str, b64: str) -> dict:
        data = base64.b64decode(b64 or "")
        aid = "A" + format(int(time.time() * 1000), "x") + "_" + uuid.uuid4().hex[:4]
        rel = f"reco_attachments/{_safe_segment(template_id)}/{aid}_{_safe_segment(filename, 120)}"
        p = self._abs(rel)
        with self._writing(p):
            _atomic_write_bytes(p, data)
        # storedPath stays relative so attachments survive a base-folder migration.
        return {"id": aid, "storedPath": rel, "size": len(data)}

    def read_attachment(self, stored_path: str) -> str | None:
        try:
            p = self._abs(stored_path)
            if not p.is_file():
                return None
            return base64.b64encode(p.read_bytes()).decode("ascii")
        except (OSError, ValueError, RuntimeError) as exc:
            log.warning("read_attachment(%s): %s", stored_path, exc)
            return None

    def delete_attachment(self, stored_path: str) -> bool:
        try:
            p = self._abs(stored_path)
            with self._writing(p):
                if p.is_file():
                    p.unlink()
                    return True
            return False
        except (OSError, ValueError, RuntimeError):
            return False

    # ---- export to a user-chosen location --------------------------------------
    def show_save_dialog(self, filename: str = "", file_types=None) -> str:
        if not self._window:
            return ""
        kwargs = {"save_filename": filename or ""}
        if file_types:
            kwargs["file_types"] = tuple(file_types) if isinstance(file_types, (list, tuple)) else (str(file_types),)
        try:
            res = self._window.create_file_dialog(_dialog_type("SAVE"), **kwargs)
        except Exception as exc:
            log.warning("save dialog failed: %s", exc)
            return ""
        if not res:
            return ""
        return str(res[0] if isinstance(res, (list, tuple)) else res)

    def write_file_abs(self, path: str, data) -> bool:
        text = data if isinstance(data, str) else json.dumps(data, indent=2, ensure_ascii=False)
        _atomic_write_text(Path(path), text)
        return True

    def write_file_b64_abs(self, path: str, b64: str) -> bool:
        _atomic_write_bytes(Path(path), base64.b64decode(b64 or ""))
        return True

    def launch_file(self, rel: str) -> bool:
        """Open a file inside ROOT with the operating system's default program."""
        p = self._abs(rel)
        if not p.exists():
            return False
        if sys.platform.startswith("win"):
            os.startfile(str(p))  # noqa: S606
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(p)])  # noqa: S603,S607
        else:
            subprocess.Popen(["xdg-open", str(p)])  # noqa: S603,S607
        return True

    # ---- generic helpers kept for other hub apps --------------------------------
    def write_csv(self, rel: str, data) -> bool:
        """data: CSV text, a list of dicts, or a JSON string of a list of dicts."""
        if isinstance(data, str):
            try:
                parsed = json.loads(data)
            except ValueError:
                return self.write_file(rel, data)
            data = parsed
        rows = [r for r in (data or []) if isinstance(r, dict)]
        header = list(dict.fromkeys(k for r in rows for k in r.keys()))
        buf = io.StringIO()
        w = csv.writer(buf, lineterminator="\r\n")
        w.writerow(header)
        for r in rows:
            w.writerow([_cell(r.get(h, "")) for h in header])
        return self.write_file(rel, buf.getvalue())

    def write_json(self, rel: str, data) -> bool:
        if isinstance(data, str):
            data = json.loads(data)
        return self.write_file(rel, json.dumps(data, indent=2, ensure_ascii=False))

    def read_file_json(self, rel: str) -> str:
        """Parsed JSON file re-serialised as a JSON string ('null' if missing/invalid)."""
        txt = self.read_file(rel)
        try:
            return json.dumps(json.loads(txt)) if txt.strip() else "null"
        except ValueError:
            return "null"

    def read_csv_json(self, rel: str) -> str:
        txt = self.read_file(rel)
        return json.dumps(list(csv.DictReader(io.StringIO(txt))) if txt else [])

    def _workbook(self, rel: str):
        try:
            import openpyxl
        except ImportError as exc:
            raise RuntimeError("openpyxl is not installed (pip install openpyxl)") from exc
        return openpyxl.load_workbook(self._abs(rel), read_only=True, data_only=True)

    def excel_sheet_names(self, rel: str) -> list[str]:
        wb = self._workbook(rel)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()

    def read_xlsx(self, rel: str, sheet: str | None = None) -> list[list]:
        wb = self._workbook(rel)
        try:
            ws = wb[sheet] if sheet else wb.worksheets[0]
            return [[("" if v is None else (v.isoformat() if hasattr(v, "isoformat") else v)) for v in row]
                    for row in ws.iter_rows(values_only=True)]
        finally:
            wb.close()

    def read_excel_json(self, rel: str, sheet: str | None = None) -> str:
        rows = self.read_xlsx(rel, sheet)
        if not rows:
            return "[]"
        header = [str(h) if h != "" else f"Column{i + 1}" for i, h in enumerate(rows[0])]
        return json.dumps([dict(zip(header, r)) for r in rows[1:]], default=str)

    def write_xlsx(self, rel: str, rows, sheet: str = "Sheet1") -> bool:
        import openpyxl
        if isinstance(rows, str):
            rows = json.loads(rows)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet[:31] or "Sheet1"
        rows = rows or []
        if rows and isinstance(rows[0], dict):
            header = list(dict.fromkeys(k for r in rows for k in r.keys()))
            ws.append(header)
            for r in rows:
                ws.append([r.get(h, "") for h in header])
        else:
            for r in rows:
                ws.append(list(r))
        buf = io.BytesIO()
        wb.save(buf)
        p = self._abs(rel)
        with self._writing(p):
            _atomic_write_bytes(p, buf.getvalue())
        return True

    # Projects: ROOT/projects/<name>/ folders used by other hub tools.
    def list_projects(self) -> list[str]:
        return self.list_directories("projects")

    def list_project_files(self, name: str) -> list[dict]:
        return self.list_files(f"projects/{_safe_segment(name)}")

    def project_metadata(self, name: str) -> dict:
        base = f"projects/{_safe_segment(name)}"
        meta_txt = self.read_file(base + "/metadata.json")
        try:
            meta = json.loads(meta_txt) if meta_txt.strip() else {}
        except ValueError:
            meta = {}
        meta.setdefault("name", name)
        meta["files"] = self.list_files(base)
        return meta

    def load_project(self, name: str) -> dict:
        base = f"projects/{_safe_segment(name)}"
        out = {"name": name, "files": {}}
        for f in self.list_files(base):
            if f["ext"] in ("json", "csv", "txt", "md", "html"):
                out["files"][f["name"]] = self.read_file(f"{base}/{f['name']}")
        return out

    def load_projects_pair(self, name_a: str, name_b: str) -> dict:
        return {"a": self.load_project(name_a), "b": self.load_project(name_b)}

    # ──────────────────────────────────────────────────────────────────────
    # Excel link (Reco "Edit in Excel" / "Load back"); the work is in excel_link.py
    # ──────────────────────────────────────────────────────────────────────
    def _xlink(self):
        if self._xl is None:
            import excel_link
            self._xl = excel_link.ExcelLinkEngine(
                abs_resolver=self._abs, root_getter=self._root, user_getter=self.current_username,
                window_getter=lambda: self._window, writing_ctx=self._writing,
                atomic_write=_atomic_write_bytes, safe_segment=_safe_segment, dialog_type=_dialog_type)
        return self._xl

    def _xl_call(self, method: str, *args, fallback=None):
        """Delegate to the engine; never raise (pywebview would reject the promise with a bare error)."""
        try:
            import excel_link
            eng = self._xlink()
            eng.maybe_sweep()
            return excel_link.json_safe(getattr(eng, method)(*args))
        except Exception as exc:
            log.exception("excel_link_%s failed", method)
            if fallback is not None:
                return fallback
            return {"ok": False, "code": "internal", "error": f"{type(exc).__name__}: {exc}"}

    def excel_link_capabilities(self) -> dict:
        return self._xl_call("capabilities")

    def excel_link_export_begin(self, head) -> dict:
        return self._xl_call("export_begin", head)

    def excel_link_export_rows(self, token, cells, base=None) -> dict:
        return self._xl_call("export_rows", token, cells, base)

    def excel_link_export_finish(self, token) -> dict:
        return self._xl_call("export_finish", token)

    def excel_link_job(self, job_id) -> dict:
        return self._xl_call("job", job_id)

    def excel_link_job_cancel(self, job_id) -> bool:
        return self._xl_call("job_cancel", job_id, fallback=False)

    def excel_link_job_release(self, job_id) -> bool:
        return self._xl_call("job_release", job_id, fallback=False)

    def excel_link_open(self, export_id) -> dict:
        return self._xl_call("open", export_id)

    def excel_link_status(self, export_id) -> dict:
        return self._xl_call("status", export_id)

    def excel_link_read_start(self, source, opts=None) -> dict:
        return self._xl_call("read_start", source, opts)

    def excel_link_read_page(self, job_id, offset=0, limit=2000) -> dict:
        return self._xl_call("read_page", job_id, offset, limit)

    def excel_link_manifest(self, export_id, with_base=False) -> dict:
        return self._xl_call("manifest", export_id, with_base)

    def excel_link_update_manifest(self, export_id, patch) -> dict:
        return self._xl_call("update_manifest", export_id, patch)

    def excel_link_list(self, data_folder="") -> dict:
        return self._xl_call("list", data_folder)

    def excel_link_close(self, export_id, delete_workbook=True) -> dict:
        return self._xl_call("close", export_id, delete_workbook)

    def excel_link_pick_file(self) -> dict:
        return self._xl_call("pick_file")

    def excel_link_reveal(self, export_id) -> dict:
        return self._xl_call("reveal", export_id)

    def excel_link_save_copy(self, export_id) -> dict:
        return self._xl_call("save_copy", export_id)


# ──────────────────────────────────────────────────────────────────────────────
# Hub page
# ──────────────────────────────────────────────────────────────────────────────

HUB_HTML = r"""<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8"><title>Automation Hub</title>
<meta name="viewport" content="width=device-width, initial-scale=1">
<style>
:root{--bg:#f5f6f8;--panel:#fff;--text:#1f2937;--muted:#6b7280;--border:#e5e7eb;--primary:#2563eb;--primary-h:#1d4ed8;
--danger:#b91c1c;--danger-bg:#fee2e2;--ok:#047857;--ok-bg:#d1fae5;--warn:#92400e;--warn-bg:#fef3c7;--header:#1f2937}
*{box-sizing:border-box}
body{margin:0;font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,sans-serif;font-size:13px;background:var(--bg);color:var(--text)}
header{background:var(--header);color:#fff;padding:10px 18px;display:flex;align-items:center;gap:14px}
header .brand{font-weight:700;font-size:15px}
header .info{flex:1;font-size:11px;opacity:.8;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
header button{background:#374151;color:#fff;border:1px solid #4b5563}
header button:hover{background:#4b5563}
main{padding:22px;max-width:1100px;margin:0 auto}
button{border:1px solid #d1d5db;background:#fff;color:var(--text);padding:6px 12px;border-radius:4px;cursor:pointer;font:inherit;font-weight:500}
button:hover{background:#f3f4f6}
button.primary{background:var(--primary);border-color:var(--primary);color:#fff}
button.primary:hover{background:var(--primary-h)}
button.danger{color:var(--danger);border-color:#fca5a5}
button:disabled{opacity:.5;cursor:not-allowed}
input[type=text],input[type=password],select{border:1px solid #d1d5db;border-radius:4px;padding:6px 8px;font:inherit;width:100%;background:#fff;color:var(--text)}
.panel{background:var(--panel);border:1px solid var(--border);border-radius:8px;padding:18px;margin-bottom:16px}
.panel h2{margin:0 0 10px;font-size:15px}
.panel h3{margin:16px 0 8px;font-size:13px;border-bottom:1px solid var(--border);padding-bottom:4px}
.muted{color:var(--muted)}
.row{display:flex;gap:8px;align-items:center}
.row>input{flex:1}
.grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(180px,1fr));gap:14px}
.tile{background:var(--panel);border:1px solid var(--border);border-radius:10px;padding:18px;cursor:pointer;text-align:center;transition:box-shadow .15s,border-color .15s}
.tile:hover{border-color:var(--primary);box-shadow:0 4px 14px rgba(37,99,235,.15)}
.tile .ic{width:56px;height:56px;margin:0 auto 10px;border-radius:12px;background:#eef2ff;display:flex;align-items:center;justify-content:center;font-size:26px;font-weight:700;color:#3730a3;overflow:hidden}
.tile .ic img{width:56px;height:56px;object-fit:contain}
.tile .nm{font-weight:600}
.note{padding:10px 12px;border-radius:6px;margin:10px 0;font-size:12px;line-height:1.5}
.note.err{background:var(--danger-bg);color:var(--danger)}
.note.ok{background:var(--ok-bg);color:var(--ok)}
.note.warn{background:var(--warn-bg);color:var(--warn)}
.note.info{background:#eff6ff;color:#1e40af}
code{background:#f3f4f6;padding:1px 5px;border-radius:3px;font-size:12px;word-break:break-all}
.kv{display:grid;grid-template-columns:140px 1fr;gap:4px 10px;font-size:12px;margin:8px 0}
.kv b{color:var(--muted);font-weight:600}
.bar{height:12px;background:#e5e7eb;border-radius:6px;overflow:hidden;margin:10px 0}
.bar>div{height:100%;background:var(--primary);transition:width .3s}
.modal-bg{position:fixed;inset:0;background:rgba(15,23,42,.55);display:none;align-items:flex-start;justify-content:center;padding:40px 16px;overflow:auto;z-index:50}
.modal-bg.show{display:flex}
.modal{background:var(--panel);border-radius:10px;width:100%;max-width:760px;box-shadow:0 20px 50px rgba(0,0,0,.3)}
.modal-h{display:flex;justify-content:space-between;align-items:center;padding:14px 18px;border-bottom:1px solid var(--border)}
.modal-h h2{margin:0;font-size:15px}
.modal-b{padding:16px 18px;max-height:70vh;overflow:auto}
.modal-f{padding:12px 18px;border-top:1px solid var(--border);display:flex;gap:8px;justify-content:flex-end}
label.chk{display:flex;gap:8px;align-items:flex-start;margin:8px 0;cursor:pointer}
label.chk input{margin-top:2px}
ul.errs{margin:6px 0;padding-left:18px;max-height:160px;overflow:auto;font-size:12px}
.tabs{display:flex;gap:4px;margin-bottom:12px;border-bottom:1px solid var(--border)}
.tabs div{padding:8px 14px;cursor:pointer;border-bottom:2px solid transparent;color:var(--muted);font-weight:600}
.tabs div.active{color:var(--primary);border-bottom-color:var(--primary)}
.empty{padding:28px;text-align:center;color:var(--muted);border:1px dashed #d1d5db;border-radius:8px}
</style></head>
<body>
<header>
  <div class="brand">⚙ Automation Hub</div>
  <div class="info" id="hdrInfo"></div>
  <button id="btnUser" style="display:none" onclick="Hub.switchUser()">Switch user</button>
  <button onclick="Hub.openSettings()">📁 Base folder…</button>
</header>
<main id="main"><div class="empty">Loading…</div></main>

<div class="modal-bg" id="modal"><div class="modal">
  <div class="modal-h"><h2 id="mTitle"></h2><button onclick="Hub.closeModal()">✕</button></div>
  <div class="modal-b" id="mBody"></div>
  <div class="modal-f" id="mFoot"></div>
</div></div>

<script>
const Hub = {
  state: null,
  adminPwd: null,
  tab: 'switch',
  pre: null,
  pollTimer: null,

  api() { return window.pywebview && window.pywebview.api; },
  esc(v) { return v == null ? '' : String(v).replace(/[&<>"']/g, c => ({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c])); },
  size(n) { n = +n || 0; const u = ['B','KB','MB','GB','TB']; let i = 0; while (n >= 1024 && i < u.length - 1) { n /= 1024; i++; } return (i ? n.toFixed(1) : n) + ' ' + u[i]; },
  $(id) { return document.getElementById(id); },

  async boot() {
    if (Hub._booted) return;
    Hub._booted = true;
    await Hub.refresh();
    const m = Hub.state && Hub.state.migration;
    if (m && !m.finished) { Hub.openSettings(true); }
  },
  async refresh() {
    try { Hub.state = await Hub.api().init_app(); }
    catch (e) { Hub.$('main').innerHTML = '<div class="note err">Could not start: ' + Hub.esc(e && e.message || e) + '</div>'; return; }
    Hub.render();
  },

  render() {
    const s = Hub.state;
    Hub.$('hdrInfo').textContent = (s.user ? s.user.name + ' · ' : '') + (s.folder || 'No base folder set') + ' · v' + s.version;
    Hub.$('btnUser').style.display = s.user ? '' : 'none';
    if (!s.configured) return Hub.renderFirstRun();
    if (!s.folder_ok) return Hub.renderUnreachable();
    if (s.pointer) return Hub.renderPointer();
    if (!s.users_file) return Hub.renderNoUsers();
    if (!s.user) return Hub.renderLogin();
    return Hub.renderApps();
  },

  renderFirstRun() {
    let h = '<div class="panel"><h2>Welcome — choose the shared base folder</h2>';
    h += '<p class="muted">Pick the folder that holds <code>apps/</code>, <code>master/</code>, <code>reco/</code> … (usually a network share). This is asked only once; an admin can change or migrate it later.</p>';
    h += '<div class="row"><input type="text" id="firstPath" placeholder="e.g. \\\\server\\finance\\Reco"><button onclick="Hub.browseInto(\'firstPath\')">Browse…</button><button class="primary" onclick="Hub.firstUse()">Use this folder</button></div>';
    h += '<div id="firstMsg"></div></div>';
    Hub.$('main').innerHTML = h;
  },
  async firstUse(initialize) {
    const path = Hub.$('firstPath').value.trim();
    if (!path) return;
    const r = await Hub.api().set_base_folder(path, '', !!initialize);
    if (r.ok) { Hub.state = r.state; Hub.render(); return; }
    let h = '<div class="note ' + (r.needs_init ? 'warn' : 'err') + '">' + Hub.esc(r.error) + '</div>';
    if (r.needs_init) h += '<button class="primary" onclick="Hub.firstUse(true)">Create the standard folder structure here</button>';
    Hub.$('firstMsg').innerHTML = h;
  },

  renderUnreachable() {
    const s = Hub.state;
    let h = '<div class="panel"><h2>Base folder not reachable</h2>';
    h += '<div class="note err">' + Hub.esc((s.folder_info && s.folder_info.error) || 'The folder cannot be opened.') + '</div>';
    h += '<div class="kv"><b>Folder</b><code>' + Hub.esc(s.folder) + '</code></div>';
    h += '<p class="muted">If this is a network share, check the VPN / network connection and retry. An admin can point the hub to another folder.</p>';
    h += '<button class="primary" onclick="Hub.refresh()">Retry</button> <button onclick="Hub.openSettings()">Change base folder…</button></div>';
    Hub.$('main').innerHTML = h;
  },

  renderPointer() {
    const p = Hub.state.pointer;
    let h = '<div class="panel"><h2>This base folder has moved</h2>';
    h += '<div class="note info">It was migrated on ' + Hub.esc(p.migrated_at) + ' by ' + Hub.esc(p.migrated_by) + '.</div>';
    h += '<div class="kv"><b>Old folder</b><code>' + Hub.esc(Hub.state.folder) + '</code><b>New folder</b><code>' + Hub.esc(p.new_folder) + '</code></div>';
    h += '<button class="primary" onclick="Hub.followPointer()">Switch to the new folder</button> ';
    h += '<button onclick="Hub.ignorePointer()">Stay on the old folder</button><div id="ptrMsg"></div></div>';
    Hub.$('main').innerHTML = h;
  },
  async followPointer() {
    const r = await Hub.api().follow_pointer();
    if (r.ok) { Hub.state = r.state; Hub.render(); }
    else Hub.$('ptrMsg').innerHTML = '<div class="note err">' + Hub.esc(r.error) + '</div>';
  },
  ignorePointer() { Hub.state.pointer = null; Hub.render(); },

  renderNoUsers() {
    let h = '<div class="panel"><h2>No user list found</h2>';
    h += '<div class="note warn"><code>master/user.csv</code> is missing in this base folder.</div>';
    h += '<p class="muted">An admin can create the standard folder structure and a user list containing your Windows user (<b>' + Hub.esc(Hub.state.os_user) + '</b>) with access to every app in <code>apps/</code>.</p>';
    h += '<div class="row" style="max-width:420px"><input type="password" id="initPwd" placeholder="Admin password"><button class="primary" onclick="Hub.initFolder()">Initialize</button></div><div id="initMsg"></div></div>';
    Hub.$('main').innerHTML = h;
  },
  async initFolder() {
    const r = await Hub.api().initialize_folder(Hub.$('initPwd').value);
    if (r.ok) { Hub.state = r.state; Hub.render(); }
    else Hub.$('initMsg').innerHTML = '<div class="note err">' + Hub.esc(r.error) + '</div>';
  },

  renderLogin() {
    let h = '<div class="panel" style="max-width:460px;margin:30px auto"><h2>Sign in</h2>';
    h += '<p class="muted">Your Windows user <b>' + Hub.esc(Hub.state.os_user) + '</b> was not found in <code>master/user.csv</code>. Enter your user ID:</p>';
    h += '<div class="row"><input type="text" id="loginId" placeholder="User ID or email" onkeydown="if(event.key===\'Enter\')Hub.login()"><button class="primary" onclick="Hub.login()">Sign in</button></div>';
    h += '<div id="loginMsg"></div></div>';
    Hub.$('main').innerHTML = h;
    setTimeout(() => { const i = Hub.$('loginId'); if (i) i.focus(); }, 50);
  },
  async login() {
    const id = Hub.$('loginId').value.trim();
    if (!id) return;
    const r = await Hub.api().login(id);
    if (r.ok) { Hub.state = r.state; Hub.render(); }
    else Hub.$('loginMsg').innerHTML = '<div class="note err">' + Hub.esc(r.error) + '</div>';
  },
  async switchUser() { Hub.state = await Hub.api().logout(); Hub.render(); },

  renderApps() {
    const s = Hub.state;
    let h = '<div class="panel"><h2>Welcome, ' + Hub.esc(s.user.name) + '</h2>';
    if (!s.apps.length) {
      h += '<div class="empty">No apps are enabled for your user in <code>master/user.csv</code>.</div>';
    } else {
      h += '<div class="grid">';
      s.apps.forEach(a => {
        const ic = a.icon ? '<img src="' + a.icon + '" alt="">' : Hub.esc(a.name.charAt(0).toUpperCase());
        h += '<div class="tile" data-key="' + Hub.esc(a.key) + '" onclick="Hub.open(this.dataset.key)"><div class="ic">' + ic + '</div><div class="nm">' + Hub.esc(a.name) + '</div><div class="muted" style="font-size:11px">' + Hub.esc(a.file) + '</div></div>';
      });
      h += '</div>';
    }
    if (s.hidden_apps) h += '<p class="muted" style="margin-top:12px;font-size:11px">' + s.hidden_apps + ' other app(s) in the folder are not enabled for you.</p>';
    h += '</div>';
    Hub.$('main').innerHTML = h;
  },
  async open(key) {
    const r = await Hub.api().open_app(key);
    if (r && !r.ok) alert(r.error);
  },
  async browseInto(inputId) {
    const cur = Hub.$(inputId).value.trim();
    const p = await Hub.api().browse_folder(cur);
    if (p) { Hub.$(inputId).value = p; Hub.$(inputId).dispatchEvent(new Event('change')); }
  },

  /* ---------------- Base folder settings ---------------- */
  modal(title, body, foot) {
    Hub.$('mTitle').textContent = title;
    Hub.$('mBody').innerHTML = body;
    Hub.$('mFoot').innerHTML = foot || '<button onclick="Hub.closeModal()">Close</button>';
    Hub.$('modal').classList.add('show');
  },
  closeModal() {
    const m = Hub.state && Hub.state.migration;
    if (Hub.pollTimer && m && !m.finished) {
      if (!confirm('A migration is still running. Close this window? (The copy keeps running; reopen "Base folder…" to see progress.)')) return;
    }
    clearInterval(Hub.pollTimer); Hub.pollTimer = null;
    Hub.$('modal').classList.remove('show');
    Hub.refresh();
  },

  async openSettings(resumeProgress) {
    if (resumeProgress) { Hub.adminPwd = Hub.adminPwd || ''; return Hub.showProgress(); }
    if (!Hub.state.configured) { Hub.adminPwd = ''; return Hub.renderSettings(); }
    if (Hub.adminPwd != null) return Hub.renderSettings();
    Hub.modal('Admin password required',
      '<p class="muted">Changing or migrating the base folder affects every user. Enter the admin password (from <code>pwd.txt</code> in the base folder).</p>' +
      '<input type="password" id="admPwd" onkeydown="if(event.key===\'Enter\')Hub.checkPwd()"><div id="admMsg"></div>',
      '<button onclick="Hub.closeModal()">Cancel</button><button class="primary" onclick="Hub.checkPwd()">Continue</button>');
    setTimeout(() => Hub.$('admPwd') && Hub.$('admPwd').focus(), 50);
  },
  async checkPwd() {
    const pwd = Hub.$('admPwd').value;
    if (await Hub.api().verify_admin_password(pwd)) { Hub.adminPwd = pwd; Hub.renderSettings(); }
    else Hub.$('admMsg').innerHTML = '<div class="note err">Incorrect password.</div>';
  },

  async renderSettings() {
    const fi = await Hub.api().get_folder_info();
    const info = fi.info;
    let h = '<div class="kv"><b>Current folder</b><code>' + Hub.esc(fi.folder || '(not set)') + '</code>';
    if (info) {
      h += '<b>Status</b><span>' + (info.exists ? (info.writable ? '✓ reachable, writable' : '⚠ reachable, read-only') : '✗ ' + Hub.esc(info.error)) + '</span>';
      h += '<b>Contains</b><span>' + (info.markers.length ? info.markers.map(Hub.esc).join(', ') : '—') + '</span>';
    }
    h += '</div>';
    h += '<div class="tabs"><div id="tabSwitch" class="' + (Hub.tab === 'switch' ? 'active' : '') + '" onclick="Hub.setTab(\'switch\')">Switch to an existing folder</div>';
    h += '<div id="tabMigrate" class="' + (Hub.tab === 'migrate' ? 'active' : '') + '" onclick="Hub.setTab(\'migrate\')">Migrate all files to a new folder</div></div>';
    h += '<div id="tabBody"></div>';
    Hub.modal('Base folder', h, '<button onclick="Hub.closeModal()">Close</button>');
    Hub._recent = fi.recent_folders || [];
    Hub.renderTab();
  },
  setTab(t) { Hub.tab = t; Hub.pre = null; Hub.$('tabSwitch').className = t === 'switch' ? 'active' : ''; Hub.$('tabMigrate').className = t === 'migrate' ? 'active' : ''; Hub.renderTab(); },

  renderTab() {
    let h = '';
    if (Hub.tab === 'switch') {
      h += '<p class="muted">Point the hub at another folder that <b>already contains</b> the data. No files are copied. Use this after someone else has moved the share, or to open a test copy.</p>';
      h += '<div class="row"><input type="text" id="swPath" placeholder="Folder path" onchange="Hub.inspect(\'swPath\',\'swInfo\')"><button onclick="Hub.browseInto(\'swPath\')">Browse…</button></div>';
      h += '<div id="swInfo"></div>';
      if ((Hub._recent || []).length) {
        h += '<h3>Recent folders</h3>';
        Hub._recent.forEach((f, i) => { h += '<div class="row" style="margin-bottom:4px"><code style="flex:1">' + Hub.esc(f) + '</code><button onclick="Hub.useRecent(' + i + ')">Use</button></div>'; });
      }
      Hub.$('tabBody').innerHTML = h;
      Hub.$('mFoot').innerHTML = '<button onclick="Hub.closeModal()">Close</button><button class="primary" onclick="Hub.doSwitch(false)">Switch</button>';
    } else {
      h += '<p class="muted">Copies <b>every file</b> in the current base folder (apps, masters, templates, reco data, versions, attachments, Special JVs …) to the new folder, verifies the copy, then switches the hub to it. A note is left in the old folder so other users\' launchers offer to follow.</p>';
      h += '<div class="row"><input type="text" id="mgPath" placeholder="New folder (empty or new)" onchange="Hub.pre=null;Hub.inspect(\'mgPath\',\'mgInfo\')"><button onclick="Hub.browseInto(\'mgPath\')">Browse…</button></div>';
      h += '<div id="mgInfo"></div>';
      h += '<h3>Options</h3>';
      h += '<label class="chk">If a file already exists in the new folder: <select id="mgConflict" style="width:auto" onchange="Hub.pre=null"><option value="abort">Stop — destination must not contain the same files</option><option value="skip">Keep the existing file</option><option value="overwrite">Overwrite with the current file</option></select></label>';
      h += '<label class="chk"><input type="checkbox" id="mgHash"> Verify every file with a SHA-256 checksum (slower; always on when removing originals)</label>';
      h += '<label class="chk"><input type="checkbox" id="mgRemove" onchange="Hub.$(\'mgMoveBox\').style.display=this.checked?\'\':\'none\'"> Remove the files from the current folder after the copy is verified (a <b>move</b>)</label>';
      h += '<div id="mgMoveBox" style="display:none" class="note warn">Other users may still be using the current folder. Only use this when everyone has closed the apps. Type <b>MOVE</b> to confirm: <input type="text" id="mgMoveConfirm" style="width:120px"></div>';
      h += '<div id="mgPre"></div>';
      Hub.$('tabBody').innerHTML = h;
      Hub.$('mFoot').innerHTML = '<button onclick="Hub.closeModal()">Close</button><button onclick="Hub.preflight()">Check</button><button class="primary" id="mgStart" onclick="Hub.startMigration()">Start migration</button>';
    }
  },

  async inspect(inputId, outId) {
    const path = Hub.$(inputId).value.trim();
    if (!path) { Hub.$(outId).innerHTML = ''; return null; }
    const i = await Hub.api().inspect_path(path);
    let cls = 'info', msg = '';
    if (!i.exists) { cls = 'err'; msg = i.error || 'Folder not found.'; }
    else if (i.is_current) { cls = 'err'; msg = 'This is the current base folder.'; }
    else if (i.inside_current || i.contains_current) { cls = 'err'; msg = 'This folder overlaps the current base folder.'; }
    else {
      msg = (i.writable ? 'Writable' : 'Not writable') + ' · ' + (i.empty ? 'empty' : i.entries + ' item(s)') + (i.looks_like_root ? ' · looks like a base folder (' + i.markers.join(', ') + ')' : '');
      if (!i.writable) cls = 'err';
      else if (Hub.tab === 'switch' && !i.looks_like_root) { cls = 'warn'; msg += '. It has no apps/ master/ reco/ folders — switching will create the standard structure.'; }
      else if (Hub.tab === 'migrate' && !i.empty) cls = 'warn';
    }
    Hub.$(outId).innerHTML = '<div class="note ' + cls + '">' + Hub.esc(msg) + '</div>';
    return i;
  },
  useRecent(i) { Hub.$('swPath').value = Hub._recent[i]; Hub.inspect('swPath', 'swInfo'); },

  async doSwitch(initialize) {
    const path = Hub.$('swPath').value.trim();
    if (!path) return;
    const r = await Hub.api().set_base_folder(path, Hub.adminPwd || '', initialize);
    if (r.ok) {
      Hub.state = r.state;
      Hub.$('mBody').innerHTML = '<div class="note ok">Switched to <code>' + Hub.esc(path) + '</code>.' + (r.created && r.created.length ? ' Created: ' + r.created.map(Hub.esc).join(', ') : '') + '</div>';
      Hub.$('mFoot').innerHTML = '<button class="primary" onclick="Hub.closeModal()">Done</button>';
      return;
    }
    if (r.needs_init) {
      if (confirm(r.error + '\n\nCreate the standard folder structure there and switch?')) return Hub.doSwitch(true);
      return;
    }
    Hub.$('swInfo').innerHTML = '<div class="note err">' + Hub.esc(r.error) + '</div>';
  },

  migOptions() {
    return { conflict: Hub.$('mgConflict').value, verify_hash: Hub.$('mgHash').checked, remove_source: Hub.$('mgRemove').checked, switch: true };
  },
  async preflight() {
    const dest = Hub.$('mgPath').value.trim();
    if (!dest) { Hub.$('mgPre').innerHTML = '<div class="note err">Choose a destination folder.</div>'; return null; }
    Hub.$('mgPre').innerHTML = '<div class="note info">Scanning the current folder… (large shares can take a minute)</div>';
    const o = Hub.migOptions();
    const r = await Hub.api().preflight_migration(dest, Hub.adminPwd || '', o.conflict);
    Hub.pre = r;
    let h = '<div class="note ' + (r.ok ? (r.warnings.length ? 'warn' : 'ok') : 'err') + '">';
    if (r.total_files != null) h += '<b>' + r.total_files + '</b> file(s), <b>' + Hub.size(r.total_bytes) + '</b> to migrate' + (r.free_bytes != null ? ' · free space at destination: ' + Hub.size(r.free_bytes) : '') + '.';
    if (r.errors.length) h += '<ul class="errs">' + r.errors.map(e => '<li>' + Hub.esc(e) + '</li>').join('') + '</ul>';
    if (r.warnings.length) h += '<ul class="errs">' + r.warnings.map(e => '<li>' + Hub.esc(e) + '</li>').join('') + '</ul>';
    if (r.conflict_examples && r.conflict_examples.length) h += '<div class="muted">e.g. ' + r.conflict_examples.slice(0, 5).map(Hub.esc).join(', ') + '</div>';
    h += '</div>';
    Hub.$('mgPre').innerHTML = h;
    return r;
  },
  async startMigration() {
    const o = Hub.migOptions();
    if (o.remove_source && (Hub.$('mgMoveConfirm').value || '').trim().toUpperCase() !== 'MOVE') {
      alert('Type MOVE in the confirmation box to remove the originals, or untick that option.');
      return;
    }
    const pre = Hub.pre || await Hub.preflight();
    if (!pre || !pre.ok) return;
    const dest = Hub.$('mgPath').value.trim();
    if (!confirm('Migrate ' + pre.total_files + ' file(s) (' + Hub.size(pre.total_bytes) + ') to\n' + dest + '\n\nand switch the hub to it' + (o.remove_source ? ', then REMOVE the originals' : '') + '?')) return;
    const r = await Hub.api().start_migration(dest, Hub.adminPwd || '', o);
    if (!r.ok) { Hub.$('mgPre').innerHTML = '<div class="note err">' + Hub.esc(r.error) + '</div>'; return; }
    Hub.showProgress();
  },
  showProgress() {
    Hub.modal('Migrating base folder', '<div id="prog"></div>', '<button class="danger" id="btnCancelMig" onclick="Hub.cancelMigration()">Cancel</button>');
    const tick = async () => {
      const m = await Hub.api().migration_status();
      if (!m) return;
      Hub.state.migration = m;
      let h = '<div class="kv"><b>From</b><code>' + Hub.esc(m.src) + '</code><b>To</b><code>' + Hub.esc(m.dst) + '</code><b>Stage</b><span>' + Hub.esc(m.state) + '</span></div>';
      h += '<div class="bar"><div style="width:' + m.percent + '%"></div></div>';
      h += '<div class="muted">' + m.percent + '% · ' + m.done_files + ' / ' + m.total_files + ' files · ' + Hub.size(m.done_bytes) + ' / ' + Hub.size(m.total_bytes) + '</div>';
      if (m.current) h += '<div class="muted" style="font-size:11px;margin-top:4px;word-break:break-all">' + Hub.esc(m.current) + '</div>';
      if (m.message) h += '<div class="note ' + (m.state === 'done' ? 'ok' : m.state === 'failed' ? 'err' : m.state === 'cancelled' ? 'warn' : 'info') + '">' + Hub.esc(m.message) + '</div>';
      if (m.errors.length) h += '<ul class="errs">' + m.errors.map(e => '<li>' + Hub.esc(e) + '</li>').join('') + (m.error_count > m.errors.length ? '<li>… ' + (m.error_count - m.errors.length) + ' more</li>' : '') + '</ul>';
      if (m.warnings.length) h += '<ul class="errs">' + m.warnings.map(e => '<li>' + Hub.esc(e) + '</li>').join('') + '</ul>';
      Hub.$('prog').innerHTML = h;
      if (m.finished) {
        clearInterval(Hub.pollTimer); Hub.pollTimer = null;
        Hub.$('mFoot').innerHTML = '<button class="primary" onclick="Hub.closeModal()">' + (m.switched ? 'Open the new folder' : 'Close') + '</button>';
      }
    };
    clearInterval(Hub.pollTimer);
    Hub.pollTimer = setInterval(tick, 500);
    tick();
  },
  async cancelMigration() {
    if (!confirm('Cancel the migration? The base folder will not be changed.')) return;
    await Hub.api().cancel_migration();
  }
};

window.addEventListener('pywebviewready', () => Hub.boot());
(function waitForApi(n) {
  if (window.pywebview && window.pywebview.api && typeof window.pywebview.api.init_app === 'function') return Hub.boot();
  if (n > 0) setTimeout(() => waitForApi(n - 1), 100);
  else document.getElementById('main').innerHTML = '<div class="note err">The launcher bridge is not available. Start the hub with <code>python main.py</code>.</div>';
})(100);
</script>
</body></html>
"""


# ──────────────────────────────────────────────────────────────────────────────
# Entry point
# ──────────────────────────────────────────────────────────────────────────────

def main(argv: list[str] | None = None) -> None:
    argv = list(sys.argv[1:] if argv is None else argv)
    debug = "--debug" in argv
    logging.basicConfig(level=logging.DEBUG if debug else logging.INFO,
                        format="%(asctime)s %(levelname)s %(name)s: %(message)s")
    import webview

    for key, value in (("ALLOW_DOWNLOADS", True), ("ALLOW_FILE_URLS", True),
                       ("OPEN_EXTERNAL_LINKS_IN_BROWSER", True)):
        with contextlib.suppress(Exception):
            webview.settings[key] = value

    api = LauncherAPI()
    window = webview.create_window("Automation Hub", html=HUB_HTML, js_api=api,
                                   width=1360, height=860, min_size=(960, 620))
    api._set_window(window)
    with contextlib.suppress(Exception):
        window.events.closing += api._on_closing
    STORAGE_DIR.mkdir(parents=True, exist_ok=True)
    # private_mode=False keeps localStorage (theme, palette) between launches.
    webview.start(debug=debug, private_mode=False, storage_path=str(STORAGE_DIR))


if __name__ == "__main__":
    main()
