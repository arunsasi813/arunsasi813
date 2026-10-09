"""
Edit a Reco Excel-link workbook the way a user would in Excel (e2e helper).

    python tests/e2e/xl_edit.py <workbook.xlsx> <ops.json | inline JSON>

Loads the workbook with openpyxl in normal (not read-only) mode, applies the
operations and saves it in place, so the file gets a new mtime like an Excel
save. Prints a one-line JSON summary.

The data sheet is the one with a ``__reco_row_id`` header in its first 20 rows
(or the op's ``sheet``); columns are addressed by header text and data rows by
Excel row number (``row``) or by the hidden row id (``rid``).

Operations (a JSON list of objects with an ``op`` key):

  set              {row|rid, col, value, type?: date|number|text|formula, sheet?}
                   date "YYYY-MM-DD" becomes a real date; formula writes "=…";
                   without type, a cell formatted as Text (@) gets text, as in Excel
  set_many         {col, value, rids: [...], type?}
  sort             {col, reverse?}        sort the data rows (all columns)
  reverse_rows     {}                     reverse the data rows (all columns)
  shuffle_cols_from {col_index}           reverse the data rows of columns >= col_index
                                          (1-based, 4 = D) only, misaligning rows
  clear            {col, rows: N}         clear the first N data rows of a column
  dup_row          {rid}                  append a copy of that row at the bottom
  delete_row       {rid}
  add_row          {values: {header: value}}
  merge            {range}                e.g. "E5:E9"
  rename_sheet     {name}
  insert_rows_top  {n, title?}            insert n title rows above the header
  insert_blank_rows {before_rid|before_row, n}
  error            {row|rid, col, value: "#N/A"}  an error cell (t="e")
"""

from __future__ import annotations

import copy
import datetime as dt
import json
import os
import sys
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell

ROW_ID = "__reco_row_id"
SYSTEM_SHEETS = {"how to use", "lists", "_reco_meta", "_reco_base"}


class Sheet:
    """The data sheet: header row, header -> column, data rows."""

    def __init__(self, wb, name: str | None = None):
        self.wb = wb
        if name:
            self.ws = wb[name]
            self.header_row = self._find_header(self.ws) or 1
        else:
            found = [(ws, self._find_header(ws)) for ws in wb.worksheets if ws.title.lower() not in SYSTEM_SHEETS]
            found = [(ws, hr) for ws, hr in found if hr]
            if found:
                self.ws, self.header_row = found[0]
            else:
                self.ws, self.header_row = wb.worksheets[0], 1
        self.refresh()

    @staticmethod
    def _find_header(ws) -> int | None:
        for r in range(1, 21):
            for c in range(1, min(ws.max_column, 200) + 1):
                v = ws.cell(r, c).value
                if isinstance(v, str) and v.strip() == ROW_ID:
                    return r
        return None

    def refresh(self) -> None:
        self.cols = {}
        for c in range(1, self.ws.max_column + 1):
            v = self.ws.cell(self.header_row, c).value
            if v is not None and str(v).strip() and str(v).strip() not in self.cols:
                self.cols[str(v).strip()] = c
        self.rid_col = self.cols.get(ROW_ID)

    @property
    def first(self) -> int:
        return self.header_row + 1

    @property
    def last(self) -> int:
        """Last row with any value (never below the header)."""
        ws = self.ws
        for r in range(ws.max_row, self.header_row, -1):
            if any(ws.cell(r, c).value is not None for c in range(1, ws.max_column + 1)):
                return r
        return self.header_row

    @property
    def width(self) -> int:
        return self.ws.max_column

    def col(self, name) -> int:
        if isinstance(name, int):
            return name
        if name not in self.cols:
            raise SystemExit(f"xl_edit: no column {name!r} (have {list(self.cols)})")
        return self.cols[name]

    def row_of(self, op: dict) -> int:
        if op.get("row") is not None:
            return int(op["row"])
        rid = str(op.get("rid"))
        if not self.rid_col:
            raise SystemExit("xl_edit: the sheet has no __reco_row_id column")
        for r in range(self.first, self.last + 1):
            if str(self.ws.cell(r, self.rid_col).value) == rid:
                return r
        raise SystemExit(f"xl_edit: row id {rid!r} not found")

    def snapshot_row(self, r: int, cols=None) -> list:
        cols = cols or range(1, self.width + 1)
        return [(self.ws.cell(r, c).value, copy.copy(self.ws.cell(r, c)._style), self.ws.cell(r, c).data_type)
                for c in cols]

    def put_row(self, r: int, items: list, cols=None) -> None:
        cols = cols or range(1, self.width + 1)
        for c, (value, style, dtype) in zip(cols, items):
            cell = self.ws.cell(r, c)
            if isinstance(cell, MergedCell):          # Excel refuses to sort merges too
                continue
            cell.value = value
            cell._style = copy.copy(style)
            if value is not None and dtype in ("s", "e"):
                cell.data_type = dtype


def _typed(value, kind: str | None):
    if value is None:
        return None, None
    if kind == "date":
        y, m, d = (int(x) for x in str(value)[:10].split("-"))
        return dt.datetime(y, m, d), None
    if kind == "number":
        f = float(value)
        return (int(f) if f.is_integer() else f), None
    if kind == "text":
        return str(value), "s"
    if kind == "formula":
        v = str(value)
        return (v if v.startswith("=") else "=" + v), None
    return value, None


def _set(sh: Sheet, r: int, col, value, kind=None) -> None:
    cell = sh.ws.cell(r, sh.col(col))
    if kind is None and value is not None and not isinstance(value, bool) and cell.number_format == "@":
        kind = "text"                         # typing into a Text-formatted cell keeps text, as in Excel
    v, dtype = _typed(value, kind)
    cell.value = v
    if dtype:
        cell.data_type = dtype


def _sort_key(v):
    """Numbers and dates first (by value), then text (case-insensitive), then blanks."""
    if v is None:
        return (2, 0.0, "")
    if isinstance(v, dt.datetime):
        return (0, v.toordinal() + v.hour / 24 + v.minute / 1440, "")
    if isinstance(v, dt.date):
        return (0, float(v.toordinal()), "")
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return (0, float(v), "")
    return (1, 0.0, str(v).lower())


def apply(path: str, ops: list[dict]) -> dict:
    p = Path(path)
    before = p.stat().st_mtime_ns
    wb = openpyxl.load_workbook(p, keep_vba=p.suffix.lower() == ".xlsm")
    sheet_name = next((o.get("sheet") for o in ops if o.get("sheet")), None)
    sh = Sheet(wb, sheet_name)
    for op in ops:
        kind = op.get("op")
        if kind == "set":
            _set(sh, sh.row_of(op), op["col"], op.get("value"), op.get("type"))
        elif kind == "error":
            cell = sh.ws.cell(sh.row_of(op), sh.col(op["col"]))
            cell.value = op.get("value") or "#N/A"
            cell.data_type = "e"
        elif kind == "set_many":
            for rid in op.get("rids") or []:
                _set(sh, sh.row_of({"rid": rid}), op["col"], op.get("value"), op.get("type"))
        elif kind in ("sort", "reverse_rows"):
            rows = list(range(sh.first, sh.last + 1))
            data = [sh.snapshot_row(r) for r in rows]
            if kind == "sort":
                c = sh.col(op["col"]) - 1
                data.sort(key=lambda items: _sort_key(items[c][0]), reverse=bool(op.get("reverse")))
            else:
                data.reverse()
            for r, items in zip(rows, data):
                sh.put_row(r, items)
        elif kind == "shuffle_cols_from":
            cols = list(range(int(op.get("col_index") or 4), sh.width + 1))
            rows = list(range(sh.first, sh.last + 1))
            data = [sh.snapshot_row(r, cols) for r in rows]
            data.reverse()
            for r, items in zip(rows, data):
                sh.put_row(r, items, cols)
        elif kind == "clear":
            c = sh.col(op["col"])
            for r in range(sh.first, min(sh.last, sh.first + int(op.get("rows") or 0) - 1) + 1):
                sh.ws.cell(r, c).value = None
        elif kind == "dup_row":
            src = sh.row_of(op)
            items = sh.snapshot_row(src)
            sh.put_row(sh.last + 1, items)
        elif kind == "delete_row":
            sh.ws.delete_rows(sh.row_of(op), 1)
        elif kind == "add_row":
            r = sh.last + 1
            for name, value in (op.get("values") or {}).items():
                _set(sh, r, name, value)
        elif kind == "merge":
            sh.ws.merge_cells(op["range"])
        elif kind == "rename_sheet":
            sh.ws.title = op["name"]
        elif kind == "insert_rows_top":
            n = int(op.get("n") or 1)
            sh.ws.insert_rows(1, n)
            sh.header_row += n
            if op.get("title"):
                sh.ws.cell(1, 1).value = op["title"]
        elif kind == "insert_blank_rows":
            r = int(op["before_row"]) if op.get("before_row") else sh.row_of({"rid": op.get("before_rid")})
            sh.ws.insert_rows(r, int(op.get("n") or 1))
        else:
            raise SystemExit(f"xl_edit: unknown op {kind!r}")
        sh.refresh()
    wb.save(p)
    if p.stat().st_mtime_ns <= before:          # same-tick save on a coarse clock
        os.utime(p, ns=(before + 1_000_000_000, before + 1_000_000_000))
    return {"ok": True, "ops": len(ops), "sheet": sh.ws.title, "headerRow": sh.header_row,
            "rows": sh.last - sh.header_row, "file": str(p)}


def main(argv: list[str]) -> int:
    if len(argv) != 2:
        print(__doc__, file=sys.stderr)
        return 2
    arg = argv[1]
    ops = json.loads(Path(arg).read_text(encoding="utf-8") if os.path.isfile(arg) else arg)
    if isinstance(ops, dict):
        ops = [ops]
    print(json.dumps(apply(argv[0], ops)), flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
