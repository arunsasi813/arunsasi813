"""
Excel Link engine for the Reco app (``LauncherAPI.excel_link_*``).

Reco hands over a column plan plus typed cells; this module writes a real .xlsx
(openpyxl, write-only) with dropdowns, styles and hidden row ids, keeps a small
manifest for it in ROOT, watches the file for Excel saves and reads it back
(read-only, on a snapshot copy) as typed raw cells. The three-way diff runs in
JS; Python only leaves out rows whose cells are identical to what it wrote.

  ExcelLinkEngine         bridge-facing engine: export / read jobs, status,
                          manifests, listing, close + sweep, file pickers
  write_workbook(...)     build the workbook into a stream (pure)
  read_workbook(...)      read a workbook into raw cells (pure)
  raw_cell, row_sig, merged_ranges, owner_files, file_handler, snapshot,
  json_safe               small pure helpers (unit-tested without pywebview)

Every engine method returns {"ok": true, ...} or {"ok": false, "code", "error"}.
stdlib + openpyxl only; openpyxl is imported lazily so the launcher still
starts (and reports ``openpyxl: null``) without it.
"""

from __future__ import annotations

import contextlib
import copy
import datetime as _dt
import hashlib
import importlib.util
import io
import json
import logging
import math
import os
import platform
import posixpath
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import uuid
import zipfile
from collections import defaultdict
from decimal import Decimal
import functools
from functools import lru_cache
from pathlib import Path, PurePath
from typing import Any, Callable, Iterable

import folder_manager as fm

log = logging.getLogger("hub.excel_link")

API_VERSION = 1
ILLEGAL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
# Also lone surrogates (from JS strings) and the two XML non-characters.
_BAD_TEXT_RE = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff￾￿]")
_SURROGATE_RE = re.compile("[\ud800-\udfff]")
EXPORT_ID_RE = re.compile(r"^X\d{8}T\d{6}-[a-z0-9]{4}$")
LIST_NAME_RE = re.compile(r"^RecoList_\d+$")
_ISO_DATE_RE = re.compile(r"^(\d{4})-(\d{2})-(\d{2})$")
# Excel writes control characters in strings as _x000D_; openpyxl leaves them escaped.
_XML_CTRL_RE = re.compile(r"_x(00[01][0-9A-Fa-f])_")

ROW_ID = "__reco_row_id"
ROW_KEY = "__reco_key"
KINDS = ("id", "status", "text", "date", "value")
EDITORS = ("list", "combo", "text", "number", "date", "none")
PROTECT_MODES = ("soft", "locked", "none")
PATCH_KEYS = ("status", "rebased", "dismissed", "importAppend", "lastImport", "newRowSigs")
SYSTEM_SHEETS = ("how to use", "lists", "_reco_meta", "_reco_base")
EXCEL_EXTS = (".xlsx", ".xlsm")
SCRATCH_COLS = ("Notes 1 (not loaded)", "Notes 2 (not loaded)", "Notes 3 (not loaded)")

MAX_CELL_TEXT = 32767
EXCEL_MAX_ROWS = 1048576
EXCEL_MAX_COLS = 16384
CHUNK_ROWS = 5000            # rows per export_rows call
PAGE_ROWS = 2000             # rows per read_page call
META_CHUNK = 30000           # columnsJson chunk size in _reco_meta
BASE_IN_WORKBOOK_MAX = 20000
LONG_PATH_LIMIT = 200        # longer absolute paths fall back to the local area
MAX_RUNNING_JOBS = 4
JOB_TTL = 15 * 60            # finished jobs are purged after this
TOKEN_TTL = 30 * 60          # export tokens expire without an export_rows call
PICK_TTL = 60 * 60
SETTLE_SECONDS = 1.5         # same (mtime, size) for this long = save finished
CANCEL_EVERY = 1000
SNAP_DELAYS = (0.2, 0.4, 0.8, 1.6, 2.0)
SNAP_ATTEMPTS = 3
DAY = 86400.0
SWEEP_DELETE_CLOSED = 1 * DAY
SWEEP_FORGET_CLOSED = 90 * DAY
SWEEP_SNAPSHOTS = 1 * DAY
SWEEP_ORPHANS = 30 * DAY

LIMITS = {"maxRows": EXCEL_MAX_ROWS - 1, "maxCols": EXCEL_MAX_COLS,
          "chunkRows": CHUNK_ROWS, "pageRows": PAGE_ROWS, "baseInWorkbookMax": BASE_IN_WORKBOOK_MAX}

# Clock and sleep go through these so tests can move time.
_mono = time.monotonic
_now = time.time


def _sleep(seconds: float) -> None:
    time.sleep(seconds)


# ──────────────────────────────────────────────────────────────────────────────
# Errors, small helpers
# ──────────────────────────────────────────────────────────────────────────────

class LinkError(Exception):
    """An expected failure, returned to JS as {ok: false, code, error, ...}."""

    def __init__(self, code: str, message: str = "", **extra):
        super().__init__(message or code)
        self.code = code
        self.message = message or code
        self.extra = extra

    def to_dict(self) -> dict:
        return {"ok": False, "code": self.code, "error": self.message, **self.extra}


class Cancelled(Exception):
    pass


def _safe_segment(value, max_len: int = 80) -> str:
    """Same rule as main._safe_segment (main passes its own in)."""
    s = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(value or "")).strip(" .")
    return (s or "_")[:max_len]


def _atomic_write_bytes(path: Path, data: bytes) -> None:
    """Fallback for main._atomic_write_bytes (main passes its own in)."""
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.name}.{uuid.uuid4().hex[:8]}.tmp")
    try:
        with open(tmp, "wb") as fh:
            fh.write(data)
            fh.flush()
            with contextlib.suppress(OSError):
                os.fsync(fh.fileno())
        os.replace(tmp, path)
    finally:
        if tmp.exists():
            with contextlib.suppress(OSError):
                tmp.unlink()


def _dialog_type(kind: str):
    """Fallback for main._dialog_type (main passes its own in)."""
    import webview
    fd = getattr(webview, "FileDialog", None)
    if fd is not None and hasattr(fd, kind):
        return getattr(fd, kind)
    return getattr(webview, f"{kind}_DIALOG")


def _openpyxl_version() -> str | None:
    try:
        import openpyxl
        return str(openpyxl.__version__)
    except Exception:
        return None


def _as_obj(value, kind: type, what: str = "argument"):
    """pywebview passes JS objects as dicts/lists; older callers send JSON text."""
    if isinstance(value, str):
        try:
            value = json.loads(value) if value.strip() else None
        except ValueError as exc:
            raise LinkError("bad_spec", f"{what} is not valid JSON: {exc}") from exc
    if value is None:
        value = kind()
    if not isinstance(value, kind):
        raise LinkError("bad_spec", f"{what} must be a {kind.__name__}")
    return value


def _check_id(export_id) -> str:
    eid = str(export_id or "").strip()
    if not EXPORT_ID_RE.match(eid):
        raise LinkError("bad_spec", f"Not a valid export id: {export_id!r}")
    return eid


def _int(v, default: int = 0) -> int:
    try:
        return int(v)
    except (TypeError, ValueError):
        return default


def _iso_utc() -> str:
    """JS-style ISO timestamp (2026-10-06T06:42:33.120Z)."""
    d = _dt.datetime.fromtimestamp(_now(), tz=_dt.timezone.utc)
    return d.strftime("%Y-%m-%dT%H:%M:%S.") + f"{d.microsecond // 1000:03d}Z"


def _iso_local(ts: float) -> str:
    return _dt.datetime.fromtimestamp(ts).isoformat(timespec="seconds")


def _parse_iso(value) -> float | None:
    """Epoch seconds for an ISO string ('Z', offsets and naive local time)."""
    if not isinstance(value, str) or not value.strip():
        return None
    s = value.strip()
    if s.endswith(("Z", "z")):
        s = s[:-1] + "+00:00"
    try:
        d = _dt.datetime.fromisoformat(s)
    except ValueError:
        try:
            d = _dt.datetime.strptime(s[:19], "%Y-%m-%dT%H:%M:%S")
        except ValueError:
            return None
    return d.timestamp()


def _clean_text(s: str) -> tuple[str, bool]:
    """Strip characters Excel cannot store and cut at 32,767."""
    changed = False
    if _BAD_TEXT_RE.search(s):
        s = _SURROGATE_RE.sub("�", s)
        s = _BAD_TEXT_RE.sub("", s)
        changed = True
    if len(s) > MAX_CELL_TEXT:
        s = s[:MAX_CELL_TEXT]
        changed = True
    return s, changed


def _clip(s, n: int) -> str:
    s = _clean_text(str(s or ""))[0]
    return s if len(s) <= n else s[: n - 1] + "…"


def _unescape_ctrl(s: str) -> str:
    if "_x" not in s:
        return s
    return _XML_CTRL_RE.sub(lambda m: chr(int(m.group(1), 16)), s)


def _read_json(path) -> Any:
    # builtins.open on purpose (tests patch it to prove base files are not read).
    with open(path, "rb") as fh:
        return json.loads(fh.read().decode("utf-8-sig"))


def _json_bytes(obj, indent: int | None = None) -> bytes:
    sep = None if indent else (",", ":")
    return json.dumps(json_safe(obj), ensure_ascii=False, indent=indent, separators=sep).encode("utf-8")


def _sha256_file(path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _no_launch() -> bool:
    return os.environ.get("RECO_XL_NO_LAUNCH", "").strip() not in ("", "0", "false")


# ──────────────────────────────────────────────────────────────────────────────
# Pure helpers
# ──────────────────────────────────────────────────────────────────────────────

def json_safe(obj):
    """Make *obj* serialisable by json.dumps without a default= hook."""
    if obj is None or isinstance(obj, (bool, str)):
        return obj
    if isinstance(obj, int):
        return int(obj)
    if isinstance(obj, float):
        return obj if math.isfinite(obj) else None
    if isinstance(obj, dict):
        return {str(k): json_safe(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple, set, frozenset)):
        items = sorted(obj, key=str) if isinstance(obj, (set, frozenset)) else obj
        return [json_safe(v) for v in items]
    if isinstance(obj, (_dt.datetime, _dt.date, _dt.time)):
        return obj.isoformat()
    if isinstance(obj, Decimal):
        f = float(obj)
        return f if math.isfinite(f) else None
    if isinstance(obj, PurePath):
        return str(obj)
    if isinstance(obj, (bytes, bytearray)):
        return bytes(obj).decode("utf-8", "replace")
    return str(obj)


def raw_cell(value, data_type: str | None = None):
    """Typed raw cell: ['s', v] ['n', v] ['b', v] ['d', iso] ['e', code] ['f', cached, '=F'] or None.

    A formula comes out as ['f', None, '=F']; read_workbook's second pass fills
    the cached value in (or turns it into ['z', None, '=F'] when there is none).
    """
    if data_type == "f":
        f = getattr(value, "text", value)
        f = "" if f is None else str(f)
        return ["f", None, f if f.startswith("=") else "=" + f]
    if value is None:
        return None
    if isinstance(value, bool):
        return ["b", value]
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not math.isfinite(value):
            return ["e", "#NUM!"]
        return ["n", value]
    if isinstance(value, _dt.datetime):
        d = (value.replace(tzinfo=None) + _dt.timedelta(microseconds=500000)).replace(microsecond=0)
        return ["d", d.isoformat(timespec="seconds")]
    if isinstance(value, _dt.date):
        return ["d", _dt.datetime(value.year, value.month, value.day).isoformat(timespec="seconds")]
    if isinstance(value, _dt.time):
        return ["s", value.replace(microsecond=0, tzinfo=None).isoformat(timespec="seconds")]
    if isinstance(value, str):
        if data_type == "e":
            return ["e", value]
        return ["s", _unescape_ctrl(value)]
    if isinstance(value, Decimal):
        return raw_cell(float(value))
    return ["s", str(value)]


def _num_token(v) -> str:
    # 15 significant digits: openpyxl writes floats with %.16g and Excel with up
    # to 17, so the last digit is not stable across a save.
    f = float(v)
    if f == 0:
        return "0"
    return format(f, ".15g")


def row_sig(values: Iterable, width: int | None = None, formulas: Iterable[int] = ()) -> str:
    """16-hex signature of a data row, shared by the writer and the reader.

    *formulas* holds the positions whose value is a formula (data_type 'f').
    """
    vals = list(values)
    if width is not None:
        vals = vals[:width] + [None] * max(0, width - len(vals))
    fset = set(formulas or ())
    toks = []
    for i, v in enumerate(vals):
        if i in fset and v is not None:
            toks.append("f" + str(getattr(v, "text", v)))
        elif v is None or (isinstance(v, str) and v == ""):
            toks.append("\x00")
        elif isinstance(v, str):
            toks.append("s" + _unescape_ctrl(v).replace("\r\n", "\n").replace("\r", "\n"))
        elif isinstance(v, bool):
            toks.append("b1" if v else "b0")
        elif isinstance(v, (int, float, Decimal)):
            toks.append("n" + _num_token(v))
        elif isinstance(v, _dt.datetime):
            if v.time() == _dt.time(0, 0):
                toks.append("d" + v.date().isoformat())
            else:
                toks.append("d" + v.replace(tzinfo=None).isoformat())
        elif isinstance(v, _dt.date):
            toks.append("d" + v.isoformat())
        else:
            toks.append("s" + str(v))
    raw = "\x1f".join(toks).encode("utf-8", "surrogatepass")
    return hashlib.sha1(raw).hexdigest()[:16]


_MERGE_RE = re.compile(rb'<(?:\w+:)?mergeCell\s+ref="([A-Z]+\d+:[A-Z]+\d+)"')
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _sheet_part(z: zipfile.ZipFile, sheet_name: str) -> str | None:
    """Zip member of a worksheet, resolved through workbook.xml and its rels."""
    import xml.etree.ElementTree as ET
    wb_part = "xl/workbook.xml"
    with contextlib.suppress(KeyError, ET.ParseError):
        for rel in ET.fromstring(z.read("_rels/.rels")):
            if rel.get("Type", "").endswith("/officeDocument"):
                wb_part = rel.get("Target", wb_part).lstrip("/")
    rels_part = posixpath.join(posixpath.dirname(wb_part), "_rels", posixpath.basename(wb_part) + ".rels")
    rid = None
    for el in ET.fromstring(z.read(wb_part)).iter():
        if el.tag.endswith("}sheet") and el.get("name") == sheet_name:
            rid = el.get(f"{{{_REL_NS}}}id")
            break
    if not rid:
        return None
    for rel in ET.fromstring(z.read(rels_part)):
        if rel.get("Id") == rid:
            target = rel.get("Target", "")
            if target.startswith("/"):
                return target.lstrip("/")
            return posixpath.normpath(posixpath.join(posixpath.dirname(wb_part), target))
    return None


def merged_ranges(xlsx_path, sheet_name: str) -> list[str]:
    """['E5:E9', ...] for one sheet, streamed from the sheet XML (read-only openpyxl has no merges)."""
    out: list[str] = []
    with zipfile.ZipFile(xlsx_path) as z:
        part = _sheet_part(z, sheet_name)
        if not part:
            return out
        with z.open(part) as fh:
            tail = b""
            while True:
                chunk = fh.read(1 << 20)
                if not chunk:
                    break
                buf = tail + chunk
                last = 0
                for m in _MERGE_RE.finditer(buf):
                    out.append(m.group(1).decode("ascii"))
                    last = m.end()
                tail = buf[max(last, len(buf) - 256):]
    return out


def owner_files(path) -> list[Path]:
    """Excel (~$name, shortened for long names) and LibreOffice (.~lock.name#) owner files."""
    p = Path(path)
    name = p.name
    out = []
    for cand in dict.fromkeys(("~$" + name, "~$" + name[1:], "~$" + name[2:], ".~lock." + name + "#")):
        q = p.parent / cand
        with contextlib.suppress(OSError):
            if q.exists():
                out.append(q)
    return out


def _handler_kind(exe: str | None) -> str:
    n = os.path.basename(str(exe or "")).lower()
    if not n:
        return "none"
    if n == "excel.exe" or "microsoft excel" in str(exe).lower():
        return "excel"
    if "soffice" in n or "scalc" in n or "libreoffice" in n:
        return "libreoffice"
    if n == "openwith.exe":
        return "none"
    return "other"


@lru_cache(maxsize=8)
def file_handler(ext: str = ".xlsx") -> dict:
    """Which program opens *ext*: {'kind': excel|libreoffice|other|none, 'exe': str|None}."""
    ext = ext if str(ext).startswith(".") else "." + str(ext)
    try:
        if sys.platform.startswith("win"):
            import ctypes
            from ctypes import wintypes
            buf = ctypes.create_unicode_buffer(1024)
            size = wintypes.DWORD(1024)
            # ASSOCF_INIT_IGNOREUNKNOWN (0x400), ASSOCSTR_EXECUTABLE (2)
            hr = ctypes.windll.shlwapi.AssocQueryStringW(0x400, 2, ext, None, buf, ctypes.byref(size))
            exe = buf.value if hr == 0 and buf.value else None
            return {"kind": _handler_kind(exe), "exe": exe}
        if sys.platform == "darwin":
            for app, kind in (("/Applications/Microsoft Excel.app", "excel"),
                              ("/Applications/LibreOffice.app", "libreoffice")):
                if os.path.isdir(app):
                    return {"kind": kind, "exe": app}
            return {"kind": "other", "exe": None}
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        desktop = ""
        if shutil.which("xdg-mime"):
            with contextlib.suppress(Exception):
                desktop = subprocess.run(["xdg-mime", "query", "default", mime], capture_output=True,
                                         text=True, timeout=3).stdout.strip()
        if desktop:
            kind = "libreoffice" if "libreoffice" in desktop.lower() else "other"
            return {"kind": kind, "exe": shutil.which("libreoffice") or shutil.which("soffice") or desktop}
        exe = shutil.which("libreoffice") or shutil.which("soffice")
        return {"kind": "libreoffice" if exe else "none", "exe": exe}
    except Exception as exc:  # never let a probe break capabilities()
        log.debug("file_handler(%s): %s", ext, exc)
        return {"kind": "none", "exe": None}


class BusySaving(Exception):
    pass


def snapshot(src, snap_dir, retries: int = 5, name: str | None = None, info: dict | None = None) -> Path:
    """Copy *src* into *snap_dir* and check that the copy is a complete zip.

    Copy errors are retried with backoff (Excel or an AV scanner may hold the
    file for a moment); a copy that is not a zip yet (Excel mid-save) or whose
    source changed during the copy is retried after 1 s, three times, then
    BusySaving. *info* (optional) receives the source's ``stat`` at copy time.
    """
    src = Path(src)
    snap_dir = Path(snap_dir)
    snap_dir.mkdir(parents=True, exist_ok=True)
    dst = snap_dir / (name or f"{uuid.uuid4().hex[:12]}{src.suffix or '.xlsx'}")
    delays = SNAP_DELAYS[: max(0, retries)]
    for attempt in range(SNAP_ATTEMPTS):
        for i in range(len(delays) + 1):
            try:
                before = os.stat(src)
                shutil.copy2(src, dst)
                after = os.stat(src)
                break
            except FileNotFoundError:
                raise
            except OSError:
                if i >= len(delays):
                    with contextlib.suppress(OSError):
                        dst.unlink()
                    raise BusySaving(f"{src.name} is locked")
                _sleep(delays[i])
        stable = (before.st_mtime_ns, before.st_size) == (after.st_mtime_ns, after.st_size)
        if stable and zipfile.is_zipfile(dst):
            if info is not None:
                info["stat"] = after
            return dst
        if attempt < SNAP_ATTEMPTS - 1:
            _sleep(1.0)
    with contextlib.suppress(OSError):
        dst.unlink()
    raise BusySaving(f"{src.name} is being saved")


# ──────────────────────────────────────────────────────────────────────────────
# Head validation (export_begin)
# ──────────────────────────────────────────────────────────────────────────────

def _editor(col: dict) -> str:
    ed = col.get("editor")
    if ed in EDITORS:
        return ed
    if col.get("list"):
        return "combo" if col.get("strict") is False else "list"
    return {"date": "date", "value": "number"}.get(col.get("kind"), "text")


def validate_head(head: dict, safe_segment: Callable = _safe_segment) -> dict:
    """Check the export head (§2.4); raises LinkError(bad_spec|too_large)."""
    def bad(msg):
        raise LinkError("bad_spec", msg)

    if not isinstance(head, dict):
        bad("head must be an object")
    _check_id(head.get("exportId"))
    cols = head.get("columns")
    if not isinstance(cols, list) or not cols or not all(isinstance(c, dict) for c in cols):
        bad("head.columns must be a non-empty list of objects")
    if len(cols) + len(SCRATCH_COLS) > EXCEL_MAX_COLS:
        raise LinkError("too_large", f"{len(cols)} columns is more than Excel allows")
    seen = set()
    for i, c in enumerate(cols):
        h = str(c.get("header") if c.get("header") is not None else c.get("key") or "").strip()
        if not h:
            bad(f"column {i} has no header")
        if h in seen:
            bad(f"duplicate header: {h}")
        seen.add(h)
        if c.get("kind") not in KINDS:
            bad(f"column {h}: unknown kind {c.get('kind')!r}")
        if c.get("list") is not None and not LIST_NAME_RE.match(str(c.get("list"))):
            bad(f"column {h}: list name must look like RecoList_<n>")
    if str(cols[0].get("header") or "").strip() != ROW_ID:
        bad(f"column 0 must be {ROW_ID}")
    rows = head.get("rowCount")
    if not isinstance(rows, int) or isinstance(rows, bool) or rows < 0:
        bad("head.rowCount must be a non-negative integer")
    buffer = _int(head.get("newRowBuffer")) if head.get("allowNewRows") else 0
    if buffer < 0:
        bad("head.newRowBuffer must not be negative")
    if rows + buffer + 1 > EXCEL_MAX_ROWS:
        raise LinkError("too_large", f"{rows} rows is more than Excel allows")
    name = head.get("fileName")
    if not isinstance(name, str) or not name.lower().endswith(".xlsx") or safe_segment(name, 120) != name:
        bad("head.fileName must be a plain file name ending in .xlsx")
    if not isinstance(head.get("dataFolder"), str) or not head["dataFolder"].strip():
        bad("head.dataFolder is required")
    if head.get("area", "shared") not in ("shared", "local"):
        bad("head.area must be 'shared' or 'local'")
    if (head.get("protect") or "soft") not in PROTECT_MODES:
        bad("head.protect must be soft, locked or none")
    lists = head.get("lists") or {}
    if not isinstance(lists, dict):
        bad("head.lists must be an object")
    for k, v in lists.items():
        if not LIST_NAME_RE.match(str(k)):
            bad(f"list name {k!r} must look like RecoList_<n>")
        vals = v.get("values") if isinstance(v, dict) else None
        if vals is not None and not isinstance(vals, list):
            bad(f"{k}.values must be a list")
    for key in ("meta", "manifest"):
        if head.get(key) is not None and not isinstance(head.get(key), dict):
            bad(f"head.{key} must be an object")
    if head.get("readme") is not None and not isinstance(head.get("readme"), list):
        bad("head.readme must be a list of rows")
    return head


# ──────────────────────────────────────────────────────────────────────────────
# Workbook writer (§3)
# ──────────────────────────────────────────────────────────────────────────────

def _style_table(date_fmt: str, value_fmt: str) -> dict:
    """name -> (fill, number format, locked, font kwargs, wrap)."""
    white_bold = {"bold": True, "color": "FFFFFF"}
    return {
        "Reco Hdr": ("374151", "General", True, white_bold, True),
        "Reco Hdr Ed": ("166534", "General", True, white_bold, True),
        "Reco Hdr Id": ("9CA3AF", "General", True, white_bold, False),
        "Reco Id": (None, "@", True, {"color": "6B7280"}, False),
        "Reco RO Text": ("F3F4F6", "General", True, None, False),
        "Reco RO Date": ("F3F4F6", date_fmt, True, None, False),
        "Reco RO Value": ("F3F4F6", value_fmt, True, None, False),
        "Reco Ed Text": ("FFF9DB", "@", False, None, False),
        "Reco Ed Date": ("FFF9DB", date_fmt, False, None, False),
        "Reco Ed Value": ("FFF9DB", value_fmt, False, None, False),
    }


def _register_styles(wb, date_fmt: str, value_fmt: str) -> None:
    from openpyxl.styles import Alignment, Font, NamedStyle, PatternFill, Protection
    for name, (fill, fmt, locked, font, wrap) in _style_table(date_fmt, value_fmt).items():
        st = NamedStyle(name=name)
        if fill:
            st.fill = PatternFill(fill_type="solid", start_color=fill, end_color=fill)
        st.number_format = fmt
        st.protection = Protection(locked=locked)
        if font:
            st.font = Font(**font)
        if wrap:
            st.alignment = Alignment(wrap_text=True, vertical="top")
        wb.add_named_style(st)


def _sheet_title(name) -> str:
    t = re.sub(r"[\[\]:*?/\\]", "_", str(name or "")).strip().strip("'")[:31]
    if not t or t.lower() in SYSTEM_SHEETS:
        return "Reco Data"
    return t


def _text_of(v) -> str | None:
    """A cell value written as forced text (None for empty)."""
    if v is None:
        return None
    if isinstance(v, dict):
        return _text_of(v.get("s"))
    if isinstance(v, str):
        return v if v != "" else None
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return str(int(v))
    return str(v) if not isinstance(v, (list, dict)) else json.dumps(v, ensure_ascii=False)


def _cell_value(kind: str, v):
    """(value to write, written as forced text?) for one JS cell value."""
    if v is None:
        return None, False
    if kind == "date" and isinstance(v, str):
        m = _ISO_DATE_RE.match(v)
        if m:
            with contextlib.suppress(ValueError):
                d = _dt.datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
                if 1900 <= d.year <= 9999:
                    return d, False
    if kind == "value" and isinstance(v, (int, float)) and not isinstance(v, bool):
        f = float(v)
        return (f, False) if math.isfinite(f) else (None, False)
    t = _text_of(v)
    return t, t is not None


def _col_width(col: dict) -> float:
    if col.get("kind") == "id":
        return 13
    px = col.get("widthPx")
    if isinstance(px, (int, float)) and not isinstance(px, bool) and px > 0:
        w = px / 7.0
    else:
        w = len(str(col.get("header") or "")) + 2
    return round(max(8.0, min(60.0, w)), 2)


def _runs(flags: list[bool]) -> list[tuple[int, int]]:
    """Contiguous runs of True as (first, last) indexes."""
    out, start = [], None
    for i, f in enumerate(flags + [False]):
        if f and start is None:
            start = i
        elif not f and start is not None:
            out.append((start, i - 1))
            start = None
    return out


def write_workbook(head: dict, cells_rows: list, base_rows: list, out_stream,
                   progress: Callable[[int, int], None] | None = None,
                   cancel: threading.Event | None = None) -> dict:
    """Build the Reco workbook (§3) into *out_stream*.

    Returns {'rowSigs': [...], 'warnings': [...], 'includedBase': bool, 'sheet': str}.
    Raises Cancelled when *cancel* is set (checked every 1,000 rows).
    """
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation

    cols: list[dict] = head["columns"]
    ncol = len(cols)
    n = len(cells_rows)
    protect = head.get("protect") or "soft"
    buffer = max(0, _int(head.get("newRowBuffer"))) if head.get("allowNewRows") else 0
    date_fmt = str(head.get("dateFormat") or "dd-mmm-yyyy")
    value_fmt = str(head.get("valueFormat") or "#,##0.00;(#,##0.00)")
    lists_in = head.get("lists") or {}
    warnings: list[str] = []
    sanitized: set[str] = set()

    def clean(s: str, key: str) -> str:
        s2, changed = _clean_text(s)
        if changed and key not in sanitized:
            sanitized.add(key)
            warnings.append(f"sanitized:{key}")
        return s2

    wb = Workbook(write_only=True)
    _register_styles(wb, date_fmt, value_fmt)          # once, before any sheet
    wb.properties.creator = _clip(head.get("user") or "Reco", 255)
    wb.properties.title = _clip(head.get("tplName") or "", 255)

    data_title = _sheet_title(head.get("sheetName") or "Reco Data")
    ws = wb.create_sheet(data_title)
    howto = wb.create_sheet("How to use")
    lists_ws = wb.create_sheet("Lists")
    meta_ws = wb.create_sheet("_reco_meta")
    include_base = bool(head.get("includeBase")) and n <= BASE_IN_WORKBOOK_MAX
    base_ws = wb.create_sheet("_reco_base") if include_base else None

    # ---- lists: column k of 'Lists' holds RecoList_k ---------------------------
    list_names = sorted((k for k in lists_in), key=lambda k: int(k.split("_")[1]))
    list_cols: dict[str, tuple[str, int]] = {}      # name -> (letter, count)
    list_vals: list[list] = []
    for k, name in enumerate(list_names, start=1):
        spec = lists_in.get(name) or {}
        vals = [clean(t, name) for t in (_text_of(v) for v in (spec.get("values") or [])) if t is not None]
        title = clean(_text_of(spec.get("title")) or name, name)
        list_vals.append([title] + vals)
        list_cols[name] = (get_column_letter(k), len(vals))

    # ---- per-column plan ---------------------------------------------------------
    letters = [get_column_letter(i + 1) for i in range(ncol)]
    editable = [bool(c.get("editable")) and c.get("kind") not in ("id", "status") for c in cols]
    cell_styles, hdr_styles, notes = [], [], []
    for c, ed in zip(cols, editable):
        kind = c.get("kind")
        if kind == "id":
            cell_styles.append("Reco Id")
            hdr_styles.append("Reco Hdr Id")
        else:
            suffix = {"date": "Date", "value": "Value"}.get(kind, "Text")
            cell_styles.append(("Reco Ed " if ed else "Reco RO ") + suffix)
            hdr_styles.append("Reco Hdr Ed" if ed else "Reco Hdr")
        note = str(c.get("note") or "")
        if ed and _editor(c) in ("list", "combo"):
            lname = c.get("list")
            if not lname or not list_cols.get(lname, ("", 0))[1]:
                if "No options configured" not in note:
                    note = (note + "\n" if note else "") + "No options configured"
        notes.append(note)
    status_idx = next((i for i, c in enumerate(cols) if c.get("kind") == "status"), None)
    scratch = list(SCRATCH_COLS) if protect == "locked" else []
    last_letter = get_column_letter(ncol + len(scratch))

    # ---- data sheet: dimensions and panes before the first append -------------
    for i, c in enumerate(cols):
        cd = ws.column_dimensions[letters[i]]
        cd.width = _col_width(c)
        if c.get("hidden"):
            cd.hidden = True
        if buffer and editable[i]:
            cd.style = cell_styles[i]                     # buffer rows inherit the editable style
    for j, _title in enumerate(scratch):
        cd = ws.column_dimensions[get_column_letter(ncol + 1 + j)]
        cd.width = 24
        cd.protection = Protection(locked=False)
    ws.freeze_panes = "D2"
    ws.sheet_properties.tabColor = "166534"

    style_arrays: dict[str, Any] = {}

    def set_style(cell, name: str) -> None:
        # Same as `cell.style = name` without the per-cell named-style lookup.
        arr = style_arrays.get(name)
        if arr is None:
            cell.style = name
            style_arrays[name] = copy.copy(cell._style)
        else:
            cell._style = copy.copy(arr)

    def text_cell(sheet, value: str, style: str | None = None):
        cell = WriteOnlyCell(sheet, value=value)
        cell.data_type = "s"
        if style:
            set_style(cell, style)
        return cell

    header = []
    for i, c in enumerate(cols):
        h = clean(str(c.get("header") if c.get("header") is not None else c.get("key")).strip(), str(c.get("key")))
        cell = text_cell(ws, h, hdr_styles[i])
        if notes[i]:
            cell.comment = Comment(_clip(notes[i], 2000), "Reco", height=90, width=260)
        header.append(cell)
    for t in scratch:
        header.append(text_cell(ws, t, "Reco Hdr"))
    ws.append(header)

    # ---- data rows -------------------------------------------------------------
    sigs: list[str] = []
    kinds = [c.get("kind") for c in cols]
    keys = [str(c.get("key") or c.get("header")) for c in cols]
    for r_i, row in enumerate(cells_rows):
        if cancel is not None and r_i % CANCEL_EVERY == 0 and cancel.is_set():
            raise Cancelled()
        out, written = [], []
        for i in range(ncol):
            value, as_text = _cell_value(kinds[i], row[i] if i < len(row) else None)
            if as_text:
                value = clean(value, keys[i]) or None
            cell = WriteOnlyCell(ws)
            if value is not None:
                cell.value = value
                if as_text:
                    cell.data_type = "s"                 # '=…' / '#N/A' text stays text
            set_style(cell, cell_styles[i])
            out.append(cell)
            written.append(value)
        ws.append(out)
        sigs.append(row_sig(written))
        if progress is not None and r_i % 500 == 0:
            progress(r_i, n)

    # ---- validation, filter, conditional formats, protection -------------------
    last_row = n + 1 + buffer
    if n + buffer > 0:
        for i, c in enumerate(cols):
            if not editable[i]:
                continue
            ed = _editor(c)
            dv = None
            if ed in ("list", "combo"):
                lname = c.get("list")
                if lname and list_cols.get(lname, ("", 0))[1]:
                    strict = ed == "list"
                    dv = DataValidation(type="list", formula1=lname, allow_blank=True, showErrorMessage=True,
                                        errorStyle="stop" if strict else "information", errorTitle="Reco",
                                        error=_clip("Pick a value from the list." if strict
                                                    else "Not in the list — your value will be kept.", 225))
            elif ed == "date":
                dv = DataValidation(type="date", operator="between", formula1="18264", formula2="73415",
                                    allow_blank=True, showErrorMessage=True, errorStyle="warning",
                                    errorTitle="Reco", error="Enter a date, e.g. 06-Oct-2026.")
            elif ed == "number":
                dv = DataValidation(type="decimal", operator="between", formula1="-1E+15", formula2="1E+15",
                                    allow_blank=True, showErrorMessage=True, errorStyle="warning",
                                    errorTitle="Reco", error="Enter a number.")
            prompt = str(c.get("prompt") or "").strip()
            if prompt:
                if dv is None:
                    dv = DataValidation(allow_blank=True)  # "any value", only shows the input message
                dv.showInputMessage = True
                dv.promptTitle = _clip(c.get("header"), 32)
                dv.prompt = _clip(prompt, 255)
            if dv is not None:
                dv.add(f"{letters[i]}2:{letters[i]}{last_row}")
                ws.data_validations.append(dv)
    if protect == "soft" and n > 0:
        for a, b in _runs([not e for e in editable]):
            dv = DataValidation(type="custom", formula1="FALSE", showErrorMessage=True, errorStyle="stop",
                                errorTitle="Reco",
                                error="Read-only in Reco — a change here would be ignored on Load back.")
            dv.add(f"{letters[a]}2:{letters[b]}{n + 1}")
            ws.data_validations.append(dv)
    ws.auto_filter.ref = f"A1:{last_letter}{last_row}"
    if n > 0:
        last_data = letters[-1]
        if status_idx is not None:
            s = letters[status_idx]
            ws.conditional_formatting.add(
                f"A2:{last_data}{n + 1}",
                FormulaRule(formula=[f'LEFT(${s}2,6)="Offset"'], font=Font(italic=True, color="6B7280")))
            ws.conditional_formatting.add(
                f"{s}2:{s}{n + 1}",
                CellIsRule(operator="equal", formula=['"Open"'], font=Font(color="92400E"),
                           fill=PatternFill(fill_type="solid", start_color="FEF3C7", end_color="FEF3C7")))
            ws.conditional_formatting.add(
                f"{s}2:{s}{n + 1}",
                FormulaRule(formula=[f'LEFT({s}2,6)="Offset"'], font=Font(color="166534"),
                            fill=PatternFill(fill_type="solid", start_color="DCFCE7", end_color="DCFCE7")))
    if protect == "locked":
        p = ws.protection
        p.sheet = True
        p.autoFilter = p.sort = p.formatCells = p.formatColumns = p.formatRows = False
        p.selectLockedCells = p.selectUnlockedCells = False

    # ---- How to use --------------------------------------------------------------
    readme = [r if isinstance(r, list) else [r] for r in (head.get("readme") or [])]
    widths: dict[int, int] = defaultdict(int)
    for row in readme:
        for j, v in enumerate(row):
            t = _text_of(v.get("v") if isinstance(v, dict) else v) or ""
            widths[j] = max(widths[j], max((len(x) for x in t.split("\n")), default=0))
    for j, w in widths.items():
        howto.column_dimensions[get_column_letter(j + 1)].width = max(8, min(80, w + 2))
    for r_i, row in enumerate(readme):
        out = []
        for v in row:
            opts = v if isinstance(v, dict) else {}
            raw = opts.get("v") if isinstance(v, dict) else v
            if isinstance(raw, (int, float)) and not isinstance(raw, bool):
                cell = WriteOnlyCell(howto, value=raw)
            else:
                t = _text_of(raw)
                cell = text_cell(howto, clean(t, "readme")) if t is not None else WriteOnlyCell(howto)
            bold = r_i == 0 or bool(opts.get("bold"))
            if bold or opts.get("color"):
                cell.font = Font(bold=bold, size=14 if r_i == 0 else 11,
                                 color=str(opts.get("color") or "000000").lstrip("#"))
            if opts.get("fill"):
                f = str(opts["fill"]).lstrip("#")
                cell.fill = PatternFill(fill_type="solid", start_color=f, end_color=f)
            if isinstance(cell.value, str) and len(cell.value) > 80:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            out.append(cell)
        howto.append(out)
    howto.protection.sheet = True

    # ---- Lists (hidden, protected) + defined names ------------------------------
    lists_ws.sheet_state = "hidden"
    depth = max((len(v) for v in list_vals), default=0)
    for r in range(depth):
        lists_ws.append([text_cell(lists_ws, v[r]) if r < len(v) else None for v in list_vals])
    for name, (letter, count) in list_cols.items():
        if count:
            wb.defined_names[name] = DefinedName(name, attr_text=f"'Lists'!${letter}$2:${letter}${count + 1}")
    lists_ws.protection.sheet = True

    # ---- _reco_meta (veryHidden): key / value, columnsJson in 30,000-char chunks
    meta_ws.sheet_state = "veryHidden"
    for k, v in (head.get("meta") or {}).items():
        if k == "columnsJson":
            text = v if isinstance(v, str) else json.dumps(v, ensure_ascii=False)
            chunks = [text[i:i + META_CHUNK] for i in range(0, len(text), META_CHUNK)] or [""]
            for idx, ch in enumerate(chunks):
                meta_ws.append([text_cell(meta_ws, f"columnsJson.{idx}"), text_cell(meta_ws, clean(ch, "meta"))])
            continue
        t = _text_of(v)
        if t is not None and len(t) > MAX_CELL_TEXT:
            warnings.append(f"metaTruncated:{k}")
        meta_ws.append([text_cell(meta_ws, clean(str(k), "meta")),
                        text_cell(meta_ws, clean(t, "meta")) if t is not None else None])

    # ---- _reco_base (veryHidden) ---------------------------------------------------
    if base_ws is not None:
        base_ws.sheet_state = "veryHidden"
        base_cols = (head.get("manifest") or {}).get("baseCols") or head.get("baseCols") or []
        base_ws.append([text_cell(base_ws, t) for t in ["rid", "pk", "fp", "flags"] + [str(c) for c in base_cols]])
        for r_i, row in enumerate(base_rows):
            if cancel is not None and r_i % CANCEL_EVERY == 0 and cancel.is_set():
                raise Cancelled()
            out = []
            for v in row:
                t = _text_of(v)
                out.append(text_cell(base_ws, clean(t, "base")) if t is not None else None)
            base_ws.append(out)

    if cancel is not None and cancel.is_set():
        raise Cancelled()
    wb.save(out_stream)
    if progress is not None:
        progress(n, n)
    return {"rowSigs": sigs, "warnings": warnings, "includedBase": include_base, "sheet": data_title}


# ──────────────────────────────────────────────────────────────────────────────
# Workbook reader (§2.5)
# ──────────────────────────────────────────────────────────────────────────────

def _hdr_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return _unescape_ctrl(str(v)).strip()


def _id_text(v) -> str:
    """rid / pk text of a cell (integral numbers without '.0')."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    if isinstance(v, int):
        return str(v)
    return str(getattr(v, "text", v)).strip()


def _sheets(wb) -> list:
    out = []
    for ws in wb.worksheets:
        if hasattr(ws, "reset_dimensions"):
            ws.reset_dimensions()            # never trust <dimension>: read every row
        out.append(ws)
    return out


def _header_row(ws, scan_rows: int) -> int | None:
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max(1, scan_rows), values_only=True), start=1):
        for v in row:
            if isinstance(v, str) and v.strip() == ROW_ID:
                return i
    return None


def _read_meta(wb) -> dict | None:
    if "_reco_meta" not in wb.sheetnames:
        return None
    ws = wb["_reco_meta"]
    ws.reset_dimensions()
    meta: dict[str, Any] = {}
    chunks: dict[int, str] = {}
    for row in ws.iter_rows(max_col=2, values_only=True):
        if not row or row[0] is None:
            continue
        k = _hdr_text(row[0])
        v = "" if len(row) < 2 or row[1] is None else _unescape_ctrl(_id_text(row[1]) if not isinstance(row[1], str) else row[1])
        m = re.match(r"^columnsJson\.(\d+)$", k)
        if m:
            chunks[int(m.group(1))] = v
        elif k:
            meta[k] = v
    if chunks:
        text = "".join(chunks[i] for i in sorted(chunks))
        try:
            meta["columns"] = json.loads(text)
        except ValueError:
            meta["columns"] = None
            meta["columnsError"] = "columnsJson is not valid JSON"
    elif "columnsJson" in meta:
        with contextlib.suppress(ValueError):
            meta["columns"] = json.loads(meta.pop("columnsJson"))
    return meta


def _find_data_sheet(wb, meta: dict | None, scan_rows: int, want: str | None):
    """(worksheet, sheetMatch, headerRow)."""
    sheets = _sheets(wb)
    by_name = {ws.title: ws for ws in sheets}
    meta = meta or {}
    meta_hr = max(1, _int(meta.get("headerRow"), 1))
    if want:
        if want not in by_name:
            raise LinkError("bad_sheet", f"Sheet '{want}' was not found.",
                            candidates=[ws.title for ws in sheets])
        ws = by_name[want]
        return ws, "user", _header_row(ws, scan_rows) or meta_hr
    ds = meta.get("dataSheet")
    if ds in by_name:
        hr = _header_row(by_name[ds], scan_rows)
        if hr:
            return by_name[ds], "meta", hr
    found = []
    for ws in sheets:
        if ws.title == ds or ws.title.lower() in SYSTEM_SHEETS:
            continue
        hr = _header_row(ws, scan_rows)
        if hr:
            found.append((ws, hr))
    if len(found) == 1:
        return found[0][0], "header-scan", found[0][1]
    if len(found) > 1:
        raise LinkError("ambiguous_sheet", "More than one sheet has Reco row ids. Pick the sheet to load.",
                        candidates=[ws.title for ws, _ in found])
    if ds in by_name:
        return by_name[ds], "first", meta_hr
    plain = [ws for ws in sheets if ws.title.lower() not in SYSTEM_SHEETS] or sheets
    if not plain:
        raise LinkError("unreadable", "The workbook has no worksheets.")
    return plain[0], "first", 1


def _read_base_sheet(wb) -> dict | None:
    if "_reco_base" not in wb.sheetnames:
        return None
    ws = wb["_reco_base"]
    ws.reset_dimensions()
    header, rows = None, []
    for row in ws.iter_rows(values_only=True):
        vals = ["" if v is None else _unescape_ctrl(v if isinstance(v, str) else _id_text(v)) for v in row]
        if header is None:
            header = vals
            while header and header[-1] == "":
                header.pop()
            continue
        if not any(vals):
            continue
        vals = vals[: len(header)] + [""] * max(0, len(header) - len(vals))
        rows.append(vals)
    return {"header": header or [], "rows": rows}


def read_workbook(path, head: dict | None = None, base: dict | None = None, opts: dict | None = None,
                  progress: Callable[[int, int], None] | None = None,
                  cancel: threading.Event | None = None,
                  lookup: Callable[[dict], tuple] | None = None) -> dict:
    """Read a (snapshot of a) Reco workbook into raw cells (§2.5).

    *head*/*base* are the manifest head and base (with rowSigs) when known;
    otherwise *lookup(meta)* may return them from ``_reco_meta.exportId``.
    The result carries the kept rows under 'rows' (the engine pages them) and
    workbook properties under 'fileProps'. Raises LinkError (ambiguous_sheet,
    bad_sheet, unreadable) and Cancelled.
    """
    from openpyxl import load_workbook

    opts = opts or {}
    scan_rows = max(1, _int(opts.get("headerScanRows"), 20) or 20)
    elide = opts.get("elide", True) is not False
    warnings: list[str] = []
    try:
        wb = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise LinkError("unreadable", f"The file could not be read as an Excel workbook ({exc}).") from exc
    try:
        meta = _read_meta(wb)
        if meta and meta.get("columnsError"):
            warnings.append("metaColumnsInvalid")
        if head is None and lookup is not None and meta and meta.get("exportId"):
            with contextlib.suppress(Exception):
                head, base = lookup(meta)
        manifest_found = head is not None
        ws, sheet_match, header_row = _find_data_sheet(wb, meta, scan_rows, opts.get("sheet"))
        title = ws.title

        hdr_vals = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ()) or ()
        texts = [_hdr_text(v) for v in hdr_vals]
        while texts and texts[-1] == "":
            texts.pop()
        headers = [{"col": i, "text": t} for i, t in enumerate(texts)]
        counts: dict[str, int] = defaultdict(int)
        for t in texts:
            if t:
                counts[t] += 1
        dup_headers = [t for t, c in counts.items() if c > 1]
        rid_col = texts.index(ROW_ID) if ROW_ID in texts else -1
        key_col = texts.index(ROW_KEY) if ROW_KEY in texts else -1
        if rid_col < 0:
            warnings.append("noRowIdColumn")
        width = len(headers)

        sig_width = _int((head or {}).get("sigWidth"))
        if not sig_width and meta and isinstance(meta.get("columns"), list):
            sig_width = len(meta["columns"])
        sig_width = sig_width or width
        rebased = (head or {}).get("rebased") or {}
        no_elide = {str(k).split("|", 1)[0] for k in rebased} if isinstance(rebased, dict) else set()
        sig_map: dict[str, str] = {}
        if elide and base and isinstance(base.get("rowSigs"), list) and isinstance(base.get("rows"), list):
            for brow, sig in zip(base["rows"], base["rowSigs"]):
                if brow:
                    sig_map[str(brow[0])] = sig
        estimate = max(1, _int((meta or {}).get("rowCount")) or len(sig_map) or 1)
        scan_width = max(width, sig_width)

        def cells_of(row) -> tuple[list, int]:
            cells, nf = [], 0
            for i in range(width):
                if i < len(row):
                    c = row[i]
                    rc = raw_cell(c.value, c.data_type)
                    if rc is not None and rc[0] == "f":
                        nf += 1
                    cells.append(rc)
                else:
                    cells.append(None)
            return cells, nf

        # ---- pass 1 ------------------------------------------------------------
        kept: list[dict] = []
        unchanged: list[str] = []
        occurrences: dict[str, list[int]] = defaultdict(list)
        elided_at: dict[str, list[int]] = defaultdict(list)
        formula_cells = rows_read = blank = 0
        r = header_row
        for row in ws.iter_rows(min_row=header_row + 1):
            r += 1
            if cancel is not None and r % CANCEL_EVERY == 0 and cancel.is_set():
                raise Cancelled()
            vals = [c.value for c in row[:scan_width]]
            if all(v is None or v == "" for v in vals):
                blank += 1
                continue
            rows_read += 1
            if progress is not None and rows_read % 1000 == 0:
                progress(min(rows_read, estimate), estimate)
            rid = _id_text(vals[rid_col]) if 0 <= rid_col < len(vals) else ""
            pk = _id_text(vals[key_col]) if 0 <= key_col < len(vals) else ""
            if rid:
                occurrences[rid].append(r)
            if sig_map and rid and rid not in no_elide and rid in sig_map:
                fpos = [i for i, c in enumerate(row[:sig_width]) if c.data_type == "f"]
                if row_sig(vals[:sig_width], sig_width, fpos) == sig_map[rid]:
                    unchanged.append(rid)
                    elided_at[rid].append(r)
                    continue
            cells, nf = cells_of(row)
            formula_cells += nf
            kept.append({"r": r, "rid": rid, "pk": pk, "cells": cells})

        # ---- duplicate rids: collect elided occurrences in a targeted pass ------
        dup_rids = {rid: rs for rid, rs in occurrences.items() if len(rs) > 1}
        need = {r for rid in dup_rids for r in elided_at.get(rid, [])}
        if need:
            again = {rid for rid in dup_rids if rid in elided_at}
            unchanged = [x for x in unchanged if x not in again]
            r = header_row
            last = max(need)
            for row in ws.iter_rows(min_row=header_row + 1, max_row=last):
                r += 1
                if r not in need:
                    continue
                vals = [c.value for c in row]
                cells, nf = cells_of(row)
                formula_cells += nf
                kept.append({"r": r, "rid": _id_text(vals[rid_col]) if 0 <= rid_col < len(vals) else "",
                             "pk": _id_text(vals[key_col]) if 0 <= key_col < len(vals) else "",
                             "cells": cells})
            kept.sort(key=lambda k: k["r"])

        props = wb.properties
        file_props = {"lastModifiedBy": getattr(props, "lastModifiedBy", None),
                      "modified": props.modified.isoformat() if getattr(props, "modified", None) else None}
        base_wb = _read_base_sheet(wb) if not (base and base.get("rows") is not None) else None
    finally:
        wb.close()

    # ---- pass 2: cached values of formulas (kept rows only) ----------------------
    if formula_cells:
        by_r = {k["r"]: k for k in kept if any(c is not None and c[0] == "f" for c in k["cells"])}
        wb2 = load_workbook(path, read_only=True, data_only=True)
        try:
            ws2 = wb2[title]
            ws2.reset_dimensions()
            r = header_row
            for row in ws2.iter_rows(min_row=header_row + 1, max_row=max(by_r)):
                r += 1
                k = by_r.get(r)
                if k is None:
                    continue
                for i, c in enumerate(k["cells"]):
                    if c is None or c[0] != "f":
                        continue
                    cached = row[i] if i < len(row) else None
                    value = getattr(cached, "value", None)
                    if value is None:
                        k["cells"][i] = ["z", None, c[2]]
                    else:
                        k["cells"][i] = ["f", raw_cell(value, cached.data_type), c[2]]
        finally:
            wb2.close()
        for k in by_r.values():                 # rows past the end of the cached sheet
            k["cells"] = [["z", None, c[2]] if c is not None and c[0] == "f" and c[1] is None else c
                          for c in k["cells"]]

    merged = []
    with contextlib.suppress(Exception):
        merged = merged_ranges(path, title)

    return {
        "meta": meta,
        "exportId": (meta or {}).get("exportId") or None,
        "manifestFound": manifest_found,
        "baseInWorkbook": base_wb is not None,
        "base": base_wb,
        "sheet": title,
        "sheetMatch": sheet_match,
        "headerRow": header_row,
        "headers": headers,
        "dupHeaders": dup_headers,
        "merged": merged,
        "ridCol": rid_col,
        "keyCol": key_col,
        "sigWidth": sig_width,
        "rowsRead": rows_read,
        "rowsSent": len(kept),
        "unchangedRids": unchanged,
        "dupRids": dup_rids,
        "blankRows": blank,
        "formulaCells": formula_cells,
        "warnings": warnings,
        "fileProps": file_props,
        "rows": kept,
    }


# ──────────────────────────────────────────────────────────────────────────────
# Jobs
# ──────────────────────────────────────────────────────────────────────────────

class Job:
    """A background export or read; polled through ExcelLinkEngine.job()."""

    def __init__(self, kind: str):
        self.id = uuid.uuid4().hex[:12]
        self.kind = kind                 # 'export' | 'read'
        self.state = "running"           # running | done | failed | cancelled
        self.percent = 0
        self.message = ""
        self.result: dict | None = None
        self.error: str | None = None
        self.code: str | None = None
        self.extra: dict = {}
        self.created = _mono()
        self.finished: float | None = None
        self.cancel = threading.Event()
        self.rows: list | None = None    # read jobs: the kept rows, paged out by read_page
        self.thread: threading.Thread | None = None

    def progress(self, done: int, total: int, lo: int = 0, hi: int = 90, message: str | None = None) -> None:
        if total > 0:
            self.percent = max(self.percent, min(hi, lo + int((hi - lo) * done / total)))
        if message:
            self.message = message

    def to_dict(self) -> dict:
        d = {"ok": True, "id": self.id, "kind": self.kind, "state": self.state, "percent": self.percent,
             "message": self.message, "result": None if self.state == "running" else self.result,
             "error": self.error, "code": self.code}
        d.update(self.extra)
        return d


# ──────────────────────────────────────────────────────────────────────────────
# Engine
# ──────────────────────────────────────────────────────────────────────────────

def _api(fn):
    """Turn LinkError into {ok: false, code, error}; other errors reach the caller."""
    @functools.wraps(fn)
    def wrapper(self, *a, **kw):
        try:
            return fn(self, *a, **kw)
        except LinkError as exc:
            return exc.to_dict()
    return wrapper


class ExcelLinkEngine:
    """Everything behind LauncherAPI.excel_link_* (see the module docstring)."""

    def __init__(self, abs_resolver: Callable[[str], Path], root_getter: Callable[[], Path],
                 user_getter: Callable[[], str] | None = None, window_getter: Callable[[], Any] | None = None,
                 writing_ctx: Callable[[Path], Any] | None = None, work_dir: str | os.PathLike | None = None,
                 atomic_write: Callable[[Path, bytes], None] | None = None,
                 safe_segment: Callable[..., str] | None = None, dialog_type: Callable[[str], Any] | None = None):
        self._abs = abs_resolver
        self._root = root_getter
        self._user = user_getter or (lambda: "")
        self._window = window_getter or (lambda: None)
        self._writing = writing_ctx or (lambda path: contextlib.nullcontext())
        self._work_dir = Path(work_dir) if work_dir else None
        self._atomic = atomic_write or _atomic_write_bytes
        self._safe = safe_segment or _safe_segment
        self._dialog_type = dialog_type or _dialog_type
        self._lock = threading.RLock()
        self._jobs: dict[str, Job] = {}
        self._tokens: dict[str, dict] = {}
        self._picks: dict[str, dict] = {}
        self._seen: dict[str, tuple] = {}
        self._head_paths: dict[str, Path] = {}
        self._swept = False

    # ---- locations -------------------------------------------------------------
    def work_dir(self) -> Path:
        """Local area: env RECO_XL_WORK_DIR, %LOCALAPPDATA%\\AutomationHub\\ExcelLink, ~/.automation_hub/excel_link."""
        if self._work_dir:
            return self._work_dir
        env = os.environ.get("RECO_XL_WORK_DIR", "").strip()
        if env:
            return Path(env)
        local = os.environ.get("LOCALAPPDATA", "").strip()
        if local:
            return Path(local) / "AutomationHub" / "ExcelLink"
        return Path.home() / ".automation_hub" / "excel_link"

    def _snap_dir(self) -> Path:
        d = self.work_dir() / ".snap"
        try:
            d.mkdir(parents=True, exist_ok=True)
            return d
        except OSError:
            d = Path(tempfile.gettempdir()) / "reco_xl_snap"
            d.mkdir(parents=True, exist_ok=True)
            return d

    def _root_dir(self) -> Path:
        try:
            r = Path(self._root())
        except Exception as exc:
            raise LinkError("no_root", "The base folder is not set.") from exc
        if not r.is_dir():
            raise LinkError("no_root", f"The base folder is not reachable: {r}")
        return r

    def _resolve(self, rel: str) -> Path:
        try:
            return Path(self._abs(rel))
        except ValueError as exc:
            raise LinkError("bad_spec", str(exc)) from exc
        except RuntimeError as exc:
            raise LinkError("no_root", str(exc)) from exc

    def _links_rel(self, data_folder: str) -> str:
        return f"reco_excel/{self._safe(data_folder, 60)}/_links"

    def _local_path(self, name) -> Path:
        name = str(name or "")
        if not name or self._safe(name, 255) != name:
            raise LinkError("bad_spec", f"Not a plain file name: {name!r}")
        wd = self.work_dir()
        p = wd / name
        if not fm.same_path(os.path.dirname(fm.norm(str(p))), fm.norm(str(wd))):
            raise LinkError("bad_spec", "The workbook path leaves the local work folder.")
        return p

    def _link_dirs(self, data_folder: str | None = None) -> list[Path]:
        if data_folder:
            d = self._resolve(self._links_rel(data_folder))
            return [d] if d.is_dir() else []
        base = self._resolve("reco_excel")
        if not base.is_dir():
            return []
        return sorted(p / "_links" for p in base.iterdir() if (p / "_links").is_dir())

    @staticmethod
    def _head_files(d: Path) -> list[Path]:
        out = []
        with contextlib.suppress(OSError):
            for p in sorted(d.iterdir(), key=lambda x: x.name):
                n = p.name
                if (n.endswith(".json") and not n.endswith(".base.json") and not n.startswith(("~$", "."))
                        and p.is_file()):
                    out.append(p)
        return out

    def _find_head(self, export_id: str) -> Path | None:
        p = self._head_paths.get(export_id)
        if p is not None and p.is_file():
            return p
        try:
            dirs = self._link_dirs()
        except LinkError:
            return None
        for d in dirs:
            cand = d / f"{export_id}.json"
            if cand.is_file():
                self._head_paths[export_id] = cand
                return cand
        return None

    def _load_head(self, export_id, code: str = "missing") -> tuple[dict, Path]:
        eid = _check_id(export_id)
        self._root_dir()
        hp = self._find_head(eid)
        if hp is None:
            raise LinkError(code, f"No Excel link {eid} was found.")
        try:
            head = _read_json(hp)
        except (OSError, ValueError) as exc:
            raise LinkError(code, f"The link manifest could not be read: {exc}") from exc
        if not isinstance(head, dict):
            raise LinkError(code, "The link manifest is not an object.")
        return head, hp

    @staticmethod
    def _base_path(head_path: Path) -> Path:
        return head_path.with_name(head_path.name[:-5] + ".base.json")

    def _workbook_path(self, head: dict) -> tuple[Path | None, str]:
        """(path, '') or (None, 'missing' | 'other_pc' | 'bad_spec'); only through the manifest."""
        f = head.get("file") if isinstance(head.get("file"), dict) else {}
        area = f.get("area")
        try:
            if area == "shared" and f.get("rel"):
                return self._resolve(str(f["rel"])), ""
            if area == "local":
                machine = f.get("machine")
                if machine and machine != platform.node():
                    return None, "other_pc"
                return self._local_path(f.get("name")), ""
        except LinkError as exc:
            return None, exc.code
        return None, "missing"

    def _file_state(self, head: dict) -> dict:
        path, err = self._workbook_path(head)
        f = head.get("file") if isinstance(head.get("file"), dict) else {}
        st = {"area": f.get("area"), "exists": False, "otherPc": err == "other_pc", "size": None,
              "mtimeNs": None, "mtime": None, "ownerFile": False, "savedSinceExport": False,
              "newSinceLastImport": False, "_path": path, "_err": err}
        if path is None:
            return st
        try:
            s = os.stat(path)
        except OSError:
            return st
        sig = (str(s.st_mtime_ns), int(s.st_size))
        saved = sig != (str(head.get("exportMtimeNs") or ""), _int(head.get("exportSize"), -1))
        li = head.get("lastImport")
        new = saved and (not isinstance(li, dict)
                         or sig != (str(li.get("mtimeNs") or ""), _int(li.get("size"), -1)))
        st.update(exists=True, size=int(s.st_size), mtimeNs=sig[0], mtime=_iso_local(s.st_mtime),
                  ownerFile=bool(owner_files(path)), savedSinceExport=saved, newSinceLastImport=new)
        return st

    def _write_head(self, hp: Path, head: dict) -> None:
        self._atomic(hp, _json_bytes(head, indent=1))

    # ---- housekeeping ------------------------------------------------------------
    def _purge(self) -> None:
        now = _mono()
        with self._lock:
            for jid, job in list(self._jobs.items()):
                if job.finished is not None and now - job.finished > JOB_TTL:
                    del self._jobs[jid]
            for tok, t in list(self._tokens.items()):
                if now - t["last"] > TOKEN_TTL:
                    del self._tokens[tok]
            for tok, p in list(self._picks.items()):
                if now - p["created"] > PICK_TTL:
                    del self._picks[tok]

    def _new_job(self, kind: str) -> Job:
        with self._lock:
            running = sum(1 for j in self._jobs.values() if j.state == "running")
            if running >= MAX_RUNNING_JOBS:
                raise LinkError("busy", "Too many Excel jobs are running. Try again in a moment.")
            job = Job(kind)
            self._jobs[job.id] = job
            return job

    def _start(self, job: Job, fn: Callable, *args) -> None:
        def run():
            try:
                result = fn(job, *args)
                job.result = result
                job.percent = 100
                job.state = "done"
            except Cancelled:
                job.message = "Cancelled"
                job.state = "cancelled"
            except LinkError as exc:
                job.code, job.error, job.extra = exc.code, exc.message, dict(exc.extra)
                job.result = dict(exc.extra) or None
                job.state = "failed"
            except Exception as exc:  # surfaced through excel_link_job
                log.exception("excel link %s job failed", job.kind)
                job.code, job.error = "internal", f"{type(exc).__name__}: {exc}"
                job.state = "failed"
            finally:
                job.finished = _mono()
        job.thread = threading.Thread(target=run, name=f"xl-{job.kind}-{job.id}", daemon=True)
        job.thread.start()

    def maybe_sweep(self) -> None:
        """Run sweep() once per process, as soon as the base folder is reachable."""
        if self._swept:
            return
        try:
            self._root_dir()
        except LinkError:
            return
        self._swept = True
        try:
            self.sweep()
        except Exception:
            log.exception("excel link sweep failed")

    # ---- capabilities ------------------------------------------------------------
    def capabilities(self) -> dict:
        try:
            pywin32 = importlib.util.find_spec("win32com") is not None
        except (ImportError, ValueError):
            pywin32 = False
        try:
            wd = str(self.work_dir())
        except Exception:
            wd = None
        return {"ok": True, "api": API_VERSION, "openpyxl": _openpyxl_version(), "pywin32": pywin32,
                "platform": sys.platform, "machine": platform.node(), "handler": dict(file_handler(".xlsx")),
                "localWorkDir": wd, "limits": dict(LIMITS)}

    # ---- export ------------------------------------------------------------------
    def _choose_target(self, head: dict) -> tuple[str, str | None, Path, list[str]]:
        name = head["fileName"]
        warnings: list[str] = []
        if (head.get("area") or "shared") == "shared":
            user = str(head.get("user") or self._user() or "user")
            rel = f"reco_excel/{self._safe(head['dataFolder'], 60)}/{self._safe(user, 40)}/{name}"
            p = self._resolve(rel)
            if len(str(p)) > LONG_PATH_LIMIT:
                warnings.append("pathTooLong")
            else:
                try:
                    p.parent.mkdir(parents=True, exist_ok=True)
                    if not fm.is_writable(str(p.parent)):
                        raise PermissionError(str(p.parent))
                    return "shared", rel, p, warnings
                except PermissionError:
                    warnings.append("sharedNotWritable")
        wd = self.work_dir()
        wd.mkdir(parents=True, exist_ok=True)
        return "local", None, self._local_path(name), warnings

    @_api
    def export_begin(self, head) -> dict:
        head = _as_obj(head, dict, "head")
        if _openpyxl_version() is None:
            raise LinkError("no_openpyxl", "openpyxl is not installed on this PC (pip install openpyxl).")
        self._root_dir()
        self._purge()
        validate_head(head, self._safe)
        eid = head["exportId"]
        links = self._resolve(self._links_rel(head["dataFolder"]))
        if (links / f"{eid}.json").exists():
            raise LinkError("exists", f"Export {eid} already exists.")
        area, rel, target, warnings = self._choose_target(head)
        if target.exists():
            raise LinkError("exists", f"{target.name} already exists.")
        token = "T" + uuid.uuid4().hex
        with self._lock:
            self._tokens[token] = {"head": head, "cells": [], "base": [], "last": _mono(), "area": area,
                                   "rel": rel, "target": target, "links": links, "warnings": warnings}
        return {"ok": True, "token": token, "area": area, "warnings": warnings}

    def _token(self, token) -> dict:
        self._purge()
        with self._lock:
            tok = self._tokens.get(str(token or ""))
        if tok is None:
            raise LinkError("bad_token", "The export was not started or has expired.")
        return tok

    @_api
    def export_rows(self, token, cells, base=None) -> dict:
        tok = self._token(token)
        cells = _as_obj(cells, list, "cells")
        base = _as_obj(base, list, "base")
        head = tok["head"]
        ncol = len(head["columns"])
        if len(cells) > CHUNK_ROWS:
            raise LinkError("bad_spec", f"At most {CHUNK_ROWS} rows per call.")
        if len(base) != len(cells):
            raise LinkError("bad_spec", f"{len(cells)} cell rows but {len(base)} base rows.")
        base_cols = (head.get("manifest") or {}).get("baseCols")
        want_base = 4 + len(base_cols) if isinstance(base_cols, list) else None
        for i, row in enumerate(cells):
            if not isinstance(row, list) or len(row) != ncol:
                raise LinkError("bad_spec", f"Row {len(tok['cells']) + i}: expected {ncol} cells.")
        for i, row in enumerate(base):
            if not isinstance(row, list) or len(row) < 4 or (want_base is not None and len(row) != want_base):
                raise LinkError("bad_spec", f"Base row {len(tok['base']) + i}: expected "
                                            f"{want_base if want_base is not None else 'at least 4'} values.")
        if len(tok["cells"]) + len(cells) > head["rowCount"]:
            raise LinkError("bad_spec", "More rows than head.rowCount.")
        with self._lock:
            tok["cells"].extend(cells)
            tok["base"].extend(base)
            tok["last"] = _mono()
            return {"ok": True, "received": len(tok["cells"])}

    @_api
    def export_finish(self, token) -> dict:
        tok = self._token(token)
        if len(tok["cells"]) != tok["head"]["rowCount"]:
            raise LinkError("row_count_mismatch",
                            f"Received {len(tok['cells'])} rows, expected {tok['head']['rowCount']}.")
        job = self._new_job("export")
        with self._lock:
            self._tokens.pop(str(token), None)
        job.message = "Writing workbook"
        self._start(job, self._run_export, tok)
        return {"ok": True, "job_id": job.id}

    def _run_export(self, job: Job, tok: dict) -> dict:
        t0 = _now()
        head = dict(tok["head"])
        area, rel, target = tok["area"], tok["rel"], tok["target"]
        warnings = list(tok["warnings"])
        eid, name = head["exportId"], head["fileName"]
        links: Path = tok["links"]
        head_path, base_path = links / f"{eid}.json", links / f"{eid}.base.json"
        manifest_rel = f"{self._links_rel(head['dataFolder'])}/{eid}.json"
        user = str(head.get("user") or self._user() or "")
        meta = dict(head.get("meta") or {})
        meta.update({"exportId": eid, "area": area, "manifestRel": manifest_rel, "headerRow": "1",
                     "rowCount": str(len(tok["cells"]))})
        meta.setdefault("format", "1")
        meta.setdefault("dataFolder", head["dataFolder"])
        meta["dataSheet"] = _sheet_title(head.get("sheetName") or "Reco Data")   # the name write_workbook uses
        head["meta"] = meta
        head["user"] = user

        buf = io.BytesIO()
        res = write_workbook(head, tok["cells"], tok["base"], buf,
                             progress=lambda d, t: job.progress(d, t, 0, 90), cancel=job.cancel)
        warnings += res["warnings"]
        data = buf.getvalue()
        if job.cancel.is_set():
            raise Cancelled()

        job.progress(1, 1, 90, 95, "Saving workbook")
        try:
            target = self._save_new(target, data)
        except PermissionError:
            if area != "shared":
                raise
            warnings.append("sharedNotWritable")
            area, rel = "local", None
            self.work_dir().mkdir(parents=True, exist_ok=True)
            target = self._save_new(self._local_path(name), data)
        st = os.stat(target)
        sha = hashlib.sha256(data).hexdigest()

        job.progress(1, 1, 95, 99, "Writing manifest")
        try:
            if job.cancel.is_set():
                raise Cancelled()
            manifest = dict(head.get("manifest") or {})
            file_info = {"area": area, "rel": rel, "name": target.name}
            if area == "local":
                file_info["machine"] = platform.node()
            manifest.setdefault("format", 1)
            manifest.setdefault("tier", "P")
            manifest.setdefault("tplId", head.get("tplId"))
            manifest.setdefault("tplName", head.get("tplName"))
            manifest.setdefault("dataFolder", head["dataFolder"])
            manifest.setdefault("exportedAt", _iso_utc())
            manifest.setdefault("exportedBy", user)
            manifest.setdefault("protect", head.get("protect") or "soft")
            manifest.setdefault("dateFormat", head.get("dateFormat"))
            manifest.setdefault("baseCols", head.get("baseCols") or [])
            manifest.update({
                "exportId": eid, "file": file_info, "includeBase": res["includedBase"],
                "exportSha256": sha, "exportSize": int(st.st_size), "exportMtimeNs": str(st.st_mtime_ns),
                "sigWidth": len(head["columns"]), "rows": len(tok["cells"]), "status": "open", "closedAt": None,
                "rev": 1, "rebased": {}, "dismissed": {}, "imports": [], "lastImport": None, "newRowSigs": {},
            })
            base_doc = {"format": 1, "exportId": eid, "baseCols": manifest.get("baseCols") or [],
                        "rows": tok["base"], "rowSigs": res["rowSigs"]}
            with self._writing(base_path):
                self._atomic(base_path, _json_bytes(base_doc))
            with self._writing(head_path):
                self._write_head(head_path, manifest)
            self._head_paths[eid] = head_path
        except Cancelled:
            self._discard(target, base_path, head_path)
            raise
        except Exception as exc:
            self._discard(target, base_path, head_path)
            raise LinkError("manifest_write_failed", f"The link manifest could not be written: {exc}") from exc

        return {"exportId": eid, "area": area, "rel": rel, "name": target.name, "rows": len(tok["cells"]),
                "cols": len(head["columns"]), "seconds": round(_now() - t0, 2), "size": int(st.st_size),
                "sha256": sha, "mtimeNs": str(st.st_mtime_ns), "warnings": warnings,
                "manifestRel": manifest_rel, "includedBase": res["includedBase"], "sheet": res["sheet"]}

    def _save_new(self, target: Path, data: bytes) -> Path:
        """Atomic write that never overwrites an existing workbook."""
        with self._writing(target):
            if target.exists():
                raise LinkError("exists", f"{target.name} already exists.")
            self._atomic(target, data)
        return target

    def _discard(self, *paths: Path) -> None:
        for p in paths:
            with contextlib.suppress(OSError):
                with self._writing(p):
                    if p.exists():
                        p.unlink()

    # ---- jobs ----------------------------------------------------------------------
    @_api
    def job(self, job_id) -> dict:
        self._purge()
        with self._lock:
            job = self._jobs.get(str(job_id or ""))
        if job is None:
            raise LinkError("unknown_job", "No such job (it may have expired).")
        return job.to_dict()

    def job_cancel(self, job_id) -> bool:
        with self._lock:
            job = self._jobs.get(str(job_id or ""))
        if job is None or job.state != "running":
            return False
        job.cancel.set()
        return True

    def job_release(self, job_id) -> bool:
        with self._lock:
            job = self._jobs.pop(str(job_id or ""), None)
        if job is None:
            return False
        job.cancel.set()
        job.rows = None
        return True

    # ---- open / status -------------------------------------------------------------
    def _existing_workbook(self, head: dict) -> Path:
        path, err = self._workbook_path(head)
        if err == "other_pc":
            machine = (head.get("file") or {}).get("machine")
            raise LinkError("other_pc", f"This workbook was created on {machine}. Open it there, "
                                        "or copy it here and use Load from file…", machine=machine)
        if err == "bad_spec":
            raise LinkError("bad_spec", "The link points outside the allowed folders.")
        if path is None or not path.is_file():
            raise LinkError("missing", "The workbook was not found.")
        return path

    @_api
    def open(self, export_id) -> dict:
        head, _ = self._load_head(export_id)
        path = self._existing_workbook(head)
        handler = dict(file_handler(path.suffix or ".xlsx"))
        if _no_launch():
            return {"ok": True, "action": "skipped", "handler": handler}
        try:
            if sys.platform.startswith("win"):
                os.startfile(str(path))  # noqa: S606
            elif sys.platform == "darwin":
                subprocess.Popen(["open", str(path)])  # noqa: S603,S607
            else:
                if not shutil.which("xdg-open"):
                    raise LinkError("no_app", "No program is set up to open .xlsx files.", handler=handler)
                subprocess.Popen(["xdg-open", str(path)], stdout=subprocess.DEVNULL,  # noqa: S603,S607
                                 stderr=subprocess.DEVNULL, start_new_session=True)
        except OSError as exc:
            if getattr(exc, "winerror", None) == 1155:
                raise LinkError("no_app", "No program is set up to open .xlsx files.", handler=handler) from exc
            raise LinkError("launch_failed", f"Excel could not be started: {exc}", handler=handler) from exc
        return {"ok": True, "action": "launched", "handler": handler}

    @_api
    def status(self, export_id) -> dict:
        eid = _check_id(export_id)
        head, _ = self._load_head(eid, code="missing_manifest")
        st = self._file_state(head)
        settled = False
        if st["exists"]:
            sig = (st["mtimeNs"], st["size"])
            now = _mono()
            with self._lock:
                seen = self._seen.get(eid)
                if seen is None or seen[0] != sig:
                    seen = self._seen[eid] = (sig, now)
            settled = now - seen[1] >= SETTLE_SECONDS
        else:
            self._seen.pop(eid, None)
        if st["otherPc"]:
            state = "other_pc"
        elif not st["exists"]:
            state = "missing"
        elif not st["savedSinceExport"]:
            state = "not_saved"
        elif not settled:
            state = "saving"
        else:
            state = "saved"
        f = head.get("file") if isinstance(head.get("file"), dict) else {}
        out = {k: v for k, v in st.items() if not k.startswith("_")}
        out.update(ok=True, exportId=eid, settled=settled, state=state, fileName=f.get("name"),
                   machine=f.get("machine"), status=head.get("status"))
        return out

    # ---- read ------------------------------------------------------------------------
    def _lookup_by_meta(self, meta: dict) -> tuple[dict | None, dict | None]:
        eid = str(meta.get("exportId") or "")
        if not EXPORT_ID_RE.match(eid):
            return None, None
        hp = None
        if meta.get("dataFolder"):
            with contextlib.suppress(LinkError):
                cand = self._resolve(f"{self._links_rel(meta['dataFolder'])}/{eid}.json")
                hp = cand if cand.is_file() else None
        hp = hp or self._find_head(eid)
        if hp is None:
            return None, None
        head = _read_json(hp)
        if not isinstance(head, dict) or head.get("exportId", eid) != eid:
            return None, None
        return head, self._load_base(hp)

    def _load_base(self, head_path: Path | None) -> dict | None:
        if head_path is None:
            return None
        bp = self._base_path(head_path)
        if not bp.is_file():
            return None
        try:
            b = _read_json(bp)
        except (OSError, ValueError):
            return None
        return b if isinstance(b, dict) else None

    @_api
    def read_start(self, source, opts=None) -> dict:
        source = _as_obj(source, dict, "source")
        opts = _as_obj(opts, dict, "opts")
        self._purge()
        eid = None
        head = head_path = None
        if source.get("exportId"):
            head, head_path = self._load_head(source["exportId"])
            eid = head.get("exportId") or str(source["exportId"])
            path = self._existing_workbook(head)
        elif source.get("token"):
            with self._lock:
                pick = self._picks.get(str(source["token"]))
            if pick is None:
                raise LinkError("bad_token", "Pick the file again.")
            path = pick["path"]
            if not path.is_file():
                raise LinkError("missing", f"{path.name} was not found.")
        else:
            raise LinkError("bad_spec", "source needs exportId or token.")
        ext = path.suffix.lower()
        if ext not in EXCEL_EXTS:
            raise LinkError("unsupported_format", "Save as Excel Workbook (*.xlsx)", ext=ext)
        job = self._new_job("read")
        job.message = "Reading workbook"
        self._start(job, self._run_read, path, eid, head, head_path, opts)
        return {"ok": True, "job_id": job.id}

    def _run_read(self, job: Job, path: Path, eid: str | None, head: dict | None,
                  head_path: Path | None, opts: dict) -> dict:
        info: dict = {}
        try:
            snap = snapshot(path, self._snap_dir(), name=f"{job.id}{path.suffix.lower()}", info=info)
        except FileNotFoundError as exc:
            raise LinkError("missing", f"{path.name} was not found.") from exc
        except BusySaving as exc:
            raise LinkError("busy_saving", "Excel is still saving the file. Try again in a moment.") from exc
        try:
            st = info["stat"]
            sha = _sha256_file(snap)
            job.progress(1, 1, 0, 5, "Reading workbook")
            base = self._load_base(head_path) if head is not None else None
            res = read_workbook(snap, head, base, opts, progress=lambda d, t: job.progress(d, t, 5, 95),
                                cancel=job.cancel, lookup=self._lookup_by_meta if head is None else None)
        finally:
            with contextlib.suppress(OSError):
                snap.unlink()
        rows = res.pop("rows")
        props = res.pop("fileProps")
        res["file"] = {"name": path.name, "ext": path.suffix.lower(), "size": int(st.st_size),
                       "mtimeNs": str(st.st_mtime_ns), "mtime": _iso_local(st.st_mtime), "sha256": sha, **props}
        meta_eid = res.get("exportId")
        res["exportId"] = eid or meta_eid
        if eid and meta_eid and meta_eid != eid:
            res["warnings"].append("exportIdMismatch")
        job.rows = rows
        return res

    @_api
    def read_page(self, job_id, offset=0, limit=PAGE_ROWS) -> dict:
        with self._lock:
            job = self._jobs.get(str(job_id or ""))
        if job is None:
            raise LinkError("unknown_job", "No such job (it may have expired).")
        if job.kind != "read" or job.state != "done":
            raise LinkError("not_done", f"The read job is {job.state}.", state=job.state)
        rows = job.rows or []
        offset = max(0, _int(offset))
        limit = max(1, min(PAGE_ROWS, _int(limit, PAGE_ROWS) or PAGE_ROWS))
        end = offset + limit
        return {"ok": True, "rows": rows[offset:end], "next": end if end < len(rows) else None,
                "total": len(rows)}

    # ---- manifests --------------------------------------------------------------------
    @_api
    def manifest(self, export_id, with_base=False) -> dict:
        head, hp = self._load_head(export_id)
        out = {"ok": True, "head": head}
        if with_base:
            b = self._load_base(hp)
            out["base"] = None if b is None else {
                "baseCols": b.get("baseCols") or head.get("baseCols") or [], "rows": b.get("rows") or []}
        return out

    @_api
    def update_manifest(self, export_id, patch) -> dict:
        patch = _as_obj(patch, dict, "patch")
        unknown = [k for k in patch if k not in PATCH_KEYS]
        if unknown:
            raise LinkError("bad_patch", f"Unknown manifest keys: {', '.join(map(str, unknown))}")
        if "status" in patch and patch["status"] not in ("open", "closed"):
            raise LinkError("bad_patch", "status must be 'open' or 'closed'")
        for k in ("rebased", "dismissed", "newRowSigs"):
            if k in patch and not isinstance(patch[k], dict):
                raise LinkError("bad_patch", f"{k} must be an object")
        if "importAppend" in patch and not isinstance(patch["importAppend"], dict):
            raise LinkError("bad_patch", "importAppend must be an object")
        if "lastImport" in patch and patch["lastImport"] is not None and not isinstance(patch["lastImport"], dict):
            raise LinkError("bad_patch", "lastImport must be an object or null")
        _, hp = self._load_head(export_id)
        with self._writing(hp):
            head = _read_json(hp)
            if "status" in patch:
                head["status"] = patch["status"]
                head["closedAt"] = _iso_utc() if patch["status"] == "closed" else None
            for k in ("rebased", "dismissed", "newRowSigs"):
                if k in patch:
                    cur = head.get(k) if isinstance(head.get(k), dict) else {}
                    cur.update(patch[k])
                    head[k] = cur
            if "importAppend" in patch:
                imports = head.get("imports") if isinstance(head.get("imports"), list) else []
                imports.append(patch["importAppend"])
                head["imports"] = imports[-50:]
            if "lastImport" in patch:
                head["lastImport"] = patch["lastImport"]
            head["rev"] = _int(head.get("rev")) + 1
            self._write_head(hp, head)
        return {"ok": True, "rev": head["rev"]}

    @_api
    def list(self, data_folder="") -> dict:
        self._root_dir()
        data_folder = str(data_folder or "").strip() or None
        self.sweep(data_folder)
        links = []
        for d in self._link_dirs(data_folder):
            for hp in self._head_files(d):
                try:
                    head = _read_json(hp)
                except (OSError, ValueError):
                    continue
                if not isinstance(head, dict):
                    continue
                eid = str(head.get("exportId") or hp.stem)
                self._head_paths.setdefault(eid, hp)
                st = self._file_state(head)
                f = head.get("file") if isinstance(head.get("file"), dict) else {}
                session = head.get("session") if isinstance(head.get("session"), dict) else {}
                links.append({
                    "exportId": eid, "tplId": head.get("tplId"), "tplName": head.get("tplName"),
                    "fileName": f.get("name"), "area": f.get("area"), "tier": head.get("tier"),
                    "machine": f.get("machine"), "exportedAt": head.get("exportedAt"),
                    "exportedBy": head.get("exportedBy"),
                    "rows": head.get("rows") if head.get("rows") is not None else session.get("exportedRows"),
                    "status": head.get("status") or "open", "closedAt": head.get("closedAt"),
                    "lastImport": head.get("lastImport"), "exists": st["exists"], "size": st["size"],
                    "mtime": st["mtime"], "mtimeNs": st["mtimeNs"], "ownerFile": st["ownerFile"],
                    "savedSinceExport": st["savedSinceExport"], "newSinceLastImport": st["newSinceLastImport"],
                    "otherPc": st["otherPc"],
                })
        links.sort(key=lambda x: str(x.get("exportedAt") or ""), reverse=True)
        return {"ok": True, "links": links}

    @_api
    def close(self, export_id, delete_workbook=True) -> dict:
        _, hp = self._load_head(export_id)
        with self._writing(hp):
            head = _read_json(hp)
            head["status"] = "closed"
            head["closedAt"] = _iso_utc()
            head["rev"] = _int(head.get("rev")) + 1
            self._write_head(hp, head)
        deleted = deferred = False
        if delete_workbook:
            path, err = self._workbook_path(head)
            if err == "other_pc":
                deferred = True
            elif path is not None and path.exists():
                if owner_files(path):
                    deferred = True
                else:
                    try:
                        with self._writing(path):
                            path.unlink()
                        deleted = True
                    except PermissionError:
                        deferred = True
        return {"ok": True, "deleted": deleted, "deferred": deferred}

    def sweep(self, data_folder: str | None = None) -> dict:
        """Delete closed links' workbooks after 1 day, forget closed links after 90 days,
        drop stale snapshots and orphaned local workbooks. Never touches open links."""
        counts = {"workbooks": 0, "links": 0, "snapshots": 0, "orphans": 0}
        now = _now()
        try:
            dirs = self._link_dirs(data_folder)
        except LinkError:
            dirs = []
        for d in dirs:
            for hp in self._head_files(d):
                try:
                    head = _read_json(hp)
                except (OSError, ValueError):
                    continue
                if not isinstance(head, dict) or head.get("status") != "closed":
                    continue
                closed = _parse_iso(head.get("closedAt"))
                if closed is None:
                    with contextlib.suppress(OSError):
                        closed = hp.stat().st_mtime
                if closed is None or now - closed < SWEEP_DELETE_CLOSED:
                    continue
                path, _err = self._workbook_path(head)
                if path is not None and path.exists() and not owner_files(path):
                    with contextlib.suppress(OSError):
                        path.unlink()
                        counts["workbooks"] += 1
                if now - closed >= SWEEP_FORGET_CLOSED:
                    for p in (self._base_path(hp), hp):
                        with contextlib.suppress(OSError):
                            with self._writing(p):
                                p.unlink()
                    counts["links"] += 1
                    eid = str(head.get("exportId") or "")
                    self._head_paths.pop(eid, None)
        wd = self.work_dir()
        snap = wd / ".snap"
        if snap.is_dir():
            for p in snap.iterdir():
                with contextlib.suppress(OSError):
                    if p.is_file() and now - p.stat().st_mtime >= SWEEP_SNAPSHOTS:
                        p.unlink()
                        counts["snapshots"] += 1
        if wd.is_dir():
            old = []
            for p in wd.glob("*.xlsx"):
                with contextlib.suppress(OSError):
                    if p.is_file() and now - p.stat().st_mtime >= SWEEP_ORPHANS and not owner_files(p):
                        old.append(p)
            if old:
                referenced = self._local_names()
                if referenced is not None:
                    for p in old:
                        if p.name not in referenced:
                            with contextlib.suppress(OSError):
                                p.unlink()
                                counts["orphans"] += 1
        return counts

    def _local_names(self) -> set[str] | None:
        """Names of local-area workbooks any manifest points at (None if ROOT is unreadable)."""
        try:
            self._root_dir()
            dirs = self._link_dirs()
        except LinkError:
            return None
        names = set()
        for d in dirs:
            for hp in self._head_files(d):
                with contextlib.suppress(OSError, ValueError, AttributeError):
                    f = _read_json(hp).get("file") or {}
                    if f.get("area") == "local" and f.get("name"):
                        names.add(str(f["name"]))
        return names

    # ---- files the user picks or keeps --------------------------------------------
    @_api
    def register_pick(self, path) -> dict:
        """Map an absolute path chosen in a dialog to a read-only token (JS never sees paths)."""
        p = Path(fm.norm(str(path)))
        if not p.is_file():
            raise LinkError("missing", f"{p.name} was not found.")
        token = "P" + uuid.uuid4().hex
        with self._lock:
            self._picks[token] = {"path": p, "name": p.name, "created": _mono()}
        return {"ok": True, "token": token, "name": p.name}

    @_api
    def pick_file(self) -> dict:
        win = self._window()
        if win is None:
            raise LinkError("no_window", "No launcher window to show a file dialog in.")
        try:
            res = win.create_file_dialog(self._dialog_type("OPEN"), allow_multiple=False,
                                         file_types=("Excel workbook (*.xlsx;*.xlsm)",))
        except Exception as exc:
            raise LinkError("dialog_failed", f"The file dialog failed: {exc}") from exc
        if not res:
            return {"ok": False, "cancelled": True}
        return self.register_pick(res[0] if isinstance(res, (list, tuple)) else res)

    @_api
    def reveal(self, export_id) -> dict:
        head, _ = self._load_head(export_id)
        path = self._existing_workbook(head)
        if _no_launch():
            return {"ok": True, "action": "skipped"}
        try:
            if sys.platform.startswith("win"):
                subprocess.Popen(f'explorer /select,"{path}"')  # noqa: S603,S607
            elif sys.platform == "darwin":
                subprocess.Popen(["open", "-R", str(path)])  # noqa: S603,S607
            else:
                subprocess.Popen(["xdg-open", str(path.parent)], stdout=subprocess.DEVNULL,  # noqa: S603,S607
                                 stderr=subprocess.DEVNULL, start_new_session=True)
        except OSError as exc:
            raise LinkError("launch_failed", f"The folder could not be shown: {exc}") from exc
        return {"ok": True, "action": "launched"}

    @_api
    def save_copy(self, export_id) -> dict:
        head, _ = self._load_head(export_id)
        path = self._existing_workbook(head)
        win = self._window()
        if win is None:
            raise LinkError("no_window", "No launcher window to show a save dialog in.")
        try:
            res = win.create_file_dialog(self._dialog_type("SAVE"), save_filename=path.name,
                                         file_types=("Excel workbook (*.xlsx;*.xlsm)",))
        except Exception as exc:
            raise LinkError("dialog_failed", f"The save dialog failed: {exc}") from exc
        if not res:
            return {"ok": False, "cancelled": True}
        dest = Path(str(res[0] if isinstance(res, (list, tuple)) else res))
        if not dest.suffix:
            dest = dest.with_name(dest.name + path.suffix)
        if not fm.same_path(str(dest), str(path)):
            shutil.copy2(path, dest)
        return {"ok": True, "path": str(dest)}
